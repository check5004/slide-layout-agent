"""Build 62 faithful native templates from the pinned, MIT licensed native asset.

The public source slide, all slot bounds and modifications are recorded. This
developer command is not called by render: production always reads saved PPTX.
"""
import copy
import hashlib
import json
import math
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_DATA_LABEL_POSITION
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from slide_agent.template_engine import clone_slide


def write(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def box(shape):
    return [shape.left, shape.top, shape.width, shape.height]


def frame_spec(tf, width, height):
    sizes = [r.font.size.pt for p in tf.paragraphs for r in p.runs if r.font.size]
    size = max(sizes or [10])
    width_pt = (width - tf.margin_left - tf.margin_right) / 12700
    height_pt = (height - tf.margin_top - tf.margin_bottom) / 12700
    line_spacing = tf.paragraphs[0].line_spacing
    line_pt = (line_spacing.pt if hasattr(line_spacing, "pt") else size * (line_spacing or 1.15))
    return {"font_pt": size, "max_lines": max(1, math.floor((height_pt + 0.5) / max(size, line_pt))),
            "max_units_per_line": round(max(1, width_pt / (size * 1.10)), 2),
            "overflow": "reject; split content into a new explicitly planned slide; never shrink or truncate"}


def add_native_chart(slide, rect, categories, series, stacked=False):
    data = CategoryChartData(); data.categories = categories
    for name, values in series: data.add_series(name, values)
    sh = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED if stacked else XL_CHART_TYPE.COLUMN_CLUSTERED,
                               *[Inches(v) for v in rect], data)
    c = sh.chart
    c.has_legend = False; c.has_title = False
    c.category_axis.visible = False; c.value_axis.visible = False
    c.value_axis.minimum_scale = 0
    c.value_axis.has_major_gridlines = False
    plot = c.plots[0]; plot.gap_width = 100
    plot.has_data_labels = True
    plot.data_labels.position = XL_DATA_LABEL_POSITION.CENTER if stacked else XL_DATA_LABEL_POSITION.OUTSIDE_END
    plot.data_labels.font.size = Pt(10)
    plot.data_labels.font.name = "Yu Gothic"
    plot.data_labels.font.color.rgb = RGBColor.from_string("322014")
    for i, ser in enumerate(c.series):
        ser.format.fill.solid(); ser.format.fill.fore_color.rgb = RGBColor.from_string(["322014", "5A3921", "C5A681"][i % 3])
        ser.format.line.fill.background()
        if stacked and i<2:
            labels=copy.deepcopy(plot._element.find(qn("c:dLbls")))
            for clr in labels.xpath(".//a:srgbClr"): clr.set("val","FFFFFF")
            cat=ser._element.find(qn("c:cat"))
            ser._element.insert(list(ser._element).index(cat),labels)
    return sh


def remove_ids(slide, ids):
    for s in list(slide.shapes):
        if s.shape_id in ids: s.element.getparent().remove(s.element)


def chart_upgrade(slide, key, changes):
    if key == "chart_insight":
        remove_ids(slide, [11,12,14,15,17,18,20,21])
        add_native_chart(slide, (1.25,2.65,4.6,3.5), ["ラベル1","ラベル2","ラベル3","ラベル4"], [("指標",[32,41,49,58])])
    elif key == "stacked_bar":
        remove_ids(slide, [11,12,13,14,15,16,18,19,20,21,22,23,25,26,27,28,29,30])
        add_native_chart(slide, (0.95,2.55,11.45,3.6), ["ラベル4","ラベル5","ラベル6"],
                         [("ラベル1",[40,55,72]),("ラベル2",[30,32,33]),("ラベル3",[20,22,24])], True)
    elif key == "small_multiples":
        for x, y, ids, vals, cats in [(0.63,2.7,[12,14,16],[20,35,58],["ラベル2","ラベル3","ラベル4"]),
                                    (6.917,2.7,[20,22,24],[15,28,44],["ラベル6","ラベル7","ラベル8"]),
                                    (0.63,5.245,[28,30,32],[8,14,23],["ラベル10","ラベル11","ラベル12"]),
                                    (6.917,5.245,[36,38,40],[12,22,37],["ラベル14","ラベル15","ラベル16"])]:
            remove_ids(slide, ids)
            chart = add_native_chart(slide,(x,y,5.787,1.265),cats,[("指標",vals)])
            chart.chart.value_axis.maximum_scale = 60
    else: return
    changes.append("Shape bars replaced within their original panel by native editable charts with embedded workbooks; native chart internal padding differs from HTML.")


def chart_spec(shape):
    c = shape.chart
    xy = c.chart_type == XL_CHART_TYPE.XY_SCATTER
    cats = [] if xy else [str(x.label) for x in c.plots[0].categories]
    series = []
    for s in c.series:
        xs = [float(x.text) for x in s._element.xpath(".//c:xVal//c:pt/c:v")] if xy else []
        series.append({"name": s.name, "values": list(s.values), "x_values": xs})
    domain=[c.value_axis.minimum_scale,c.value_axis.maximum_scale]
    if domain[0] is None: domain[0]=-1e9
    if domain[1] is None: domain[1]=1e9
    if xy:
        cats=[x.text for x in c._chartSpace.xpath(".//c:dLbl/c:tx/c:rich//a:t")]
    return {"shape": shape.name, "kind": "xy" if xy else "category", "category_count": len(cats), "domain":domain,
            **({"x_domain": [0,100]} if xy else {}),
            "series_count": len(series), "point_counts": [len(s["values"]) for s in series],
            "max_label_units": 18, "bounds_emu": box(shape),
            "sample": {"categories": cats, "series": series}}


def configure_aliases(slide, entry):
    """Bind decorative data labels to chart data so values cannot contradict bars."""
    aliases={}
    shapes={s.shape_id:s for s in slide.shapes}
    key=entry["layout_id"]
    if not entry["charts"]: return aliases
    chart_key=next(iter(entry["charts"]))
    def add(sid,kind,series=0,point=None):
        s=shapes[sid]
        aliases[s.name]={"chart":chart_key,"kind":kind,"series":series,"point":point,
                         **frame_spec(s.text_frame,s.width,s.height)}
        if kind in ('cagr','x_mean'):
            aliases[s.name]['number_format']={'mode':'fixed','decimal_places':0 if kind=='cagr' else 1,
                                              'rounding':'half_even','scope':'derived value only; input values remain exact'}
    if key=="scenario_lines_cagr":
        entry["charts"][chart_key]["period_policy"]={"required":True,"field":"elapsed_years","unit":"years",
            "calendar_labels":"increasing equal year intervals, matching total elapsed_years",
            "non_year_labels":"requires explicit source-backed elapsed_years; no inference from point count"}
        entry["charts"][chart_key]["sample"]["categories"]=[str(y) for y in range(2026,2032)]
        entry["charts"][chart_key]["sample"]["elapsed_years"]=5.0
        add(17,"cagr_heading")
        entry["charts"][chart_key]["sample"]["series"][0]["name"]="上振れ"
        entry["charts"][chart_key]["sample"]["series"][1]["name"]="基準"
        entry["charts"][chart_key]["sample"]["series"][2]["name"]="下振れ"
        for i in range(3):
            add(14+i,"last",i); add(19+i*2,"cagr",i); add(25+i*2,"series_name",i)
    elif key=="delta_bars_totals":
        add(13,"sum",0); add(15,"sum",1)
        for sid,i in [(14,0),(16,1),(19,0),(21,1)]: add(sid,"series_name",i)
    elif key=="chart_insight":
        for i,sid in enumerate([13,16,19,22]): add(sid,"category",point=i)
    elif key=="stacked_bar":
        for i,sid in enumerate([17,24,31]): add(sid,"category",point=i)
        for i,sid in enumerate([33,35,37]): add(sid,"series_name",i)
    elif key=="small_multiples":
        for ck,ids in zip(entry["charts"],[[13,15,17],[21,23,25],[29,31,33],[37,39,41]]):
            chart_key=ck
            for i,sid in enumerate(ids): add(sid,"category",point=i)
    elif key=="scatter_annotated":
        add(17,"x_mean")
        entry["reference_line"]={"chart":chart_key,"shape":shapes[16].name,"label":shapes[17].name,
                                 "left":shapes[14].left,"width":shapes[14].width,"label_offset":Inches(.1),"domain":[0,100]}
    return aliases


def configure_metrics(slide, entry):
    sm = {s.shape_id: s for s in slide.shapes}
    key = entry["layout_id"]
    metrics = entry["metrics"]
    def metric(label_id, binding, **kwargs):
        label = sm[label_id]
        raw = label.text.replace("−", "-").replace("%", "")
        name = f"value_{label_id:03}"
        metrics[name] = {"binding": binding, "label": label.name, "sample": float(raw),
                         "min": 0, "max": 100, "suffix": "%" if "%" in label.text else "",
                         "label_capacity":{**frame_spec(label.text_frame,label.width,label.height),"max_lines":1},**kwargs}
        return name
    if key == "dot_matrix_share":
        for label_id in [13,115,217,319]:
            metric(label_id,"dots", shapes=[sm[i].name for i in range(label_id+1,label_id+101)], integer=True)
    elif key in ("proportional_circles", "progress_bubble_matrix"):
        pairs = [(14,15),(18,19)] if key == "proportional_circles" else [(i,i+1) for i in [18,20,22,24,27,29,31,33,36,38,40,42,45,47,49,51]]
        for sid, lid in pairs:
            s = sm[sid]; sample = float(sm[lid].text.replace("%", ""))
            maximum = 100 if key == "proportional_circles" else 60
            metric(lid,"area", shape=s.name, max=maximum,
                   center=[s.left+s.width/2,s.top+s.height/2], max_diameter=s.width*math.sqrt(maximum/sample))
    elif key in ("waterfall","true_waterfall"):
        ids = [11,14,17,20,23,26] if key == "waterfall" else [11,14,17,20,23]
        names = [metric(i+1,"bridge",min=-1e6,max=1e6) for i in ids]
        for name in names[1:-1]:metrics[name]["signed"]=True
        entry["bridge"] = {"keys":names,"shapes":[sm[i].name for i in ids],
                           "labels":[sm[i+1].name for i in ids],"cumulative":key=="true_waterfall",
                           "colors":{"start":str(sm[ids[0]].fill.fore_color.rgb),"end":str(sm[ids[-1]].fill.fore_color.rgb),
                                     "increase":"5A3921","decrease":"A22727","zero":"8A7B6B"},
                           "baseline":Inches(6.1),"height":Inches(2.25 if key=="waterfall" else 3.3),
                           "label_gap":Inches(0.26 if key=="waterfall" else 0.28)}
    elif key == "gantt":
        starts, ends = [0,1,2,3,4,8,10], [1,3,3,5,6,10,12]
        step = Inches(8.573 / 12)
        for i, sid in enumerate([51,55,59,63,67,71,75]):
            start_key, end_key = f"task_{i+1}_start", f"task_{i+1}_end"
            metrics[start_key] = {"binding":"gantt_start","shape":sm[sid].name,"sample":float(starts[i]),
                                  "min":0,"max":11,"integer":True,"end_key":end_key,"left":Inches(4.17),"step":step,"gap":Inches(.08)}
            metrics[end_key] = {"binding":"gantt_end","sample":float(ends[i]),"min":1,"max":12,"integer":True}
    elif key == "calc_flow":
        groups=[]
        for ids in [[13,16,19],[25,28,31],[37,41,45]]:
            groups.append([{"shape":sm[i].name,"label":sm[i+1].name,"key":metric(i+1,"calculation",min=0.01,max=1e6)} for i in ids])
        # The source example is not arithmetically consistent. Use an explicit
        # split of A*B into base + increment; both remain editable source data.
        for i,base in enumerate([1.5,4.5,9.0]): metrics[groups[2][i]["key"]]["sample"]=base
        for i,sid in enumerate([39,43,47]):
            key=f"increment_{i+1}"
            metrics[key]={"binding":"calculation","min":0,"max":1e6,"sample":[.5,1.5,3.0][i]}
            groups[2][i].update(extra_key=key,extra_shape=sm[sid].name)
        totals=[]
        for bar in groups[2]:
            original=sm[int(bar["shape"].split(":")[1])]
            total=slide.shapes.add_textbox(original.left-Inches(.1),Inches(3),original.width+Inches(.2),Inches(.24))
            total.name=f"static:total_{total.shape_id}"
            tf=total.text_frame;tf.auto_size=MSO_AUTO_SIZE.NONE
            tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
            run=tf.paragraphs[0].add_run();run.text="合計";run.font.name="Yu Gothic";run.font.size=Pt(10)
            run.font.color.rgb=RGBColor.from_string("322014")
            totals.append(total.name)
        entry["calculation"]={"groups":groups,"totals":totals,"baseline":Inches(5.84),"height":Inches(2.5)}


def configure_states(slide,entry):
    sm={s.shape_id:s for s in slide.shapes}; states={}; fixed=[]
    key=entry["layout_id"]
    if key=="status_heatmap_comment":
        palette={"大きく改善":"322014","改善":"5A3921","横ばい":"E2DCD2","悪化":"C5A681"}
        for sid in range(22,30):
            color=str(sm[sid].fill.fore_color.rgb)
            states[f"state_{sid:03}"]={"shape":sm[sid].name,"sample":next(k for k,v in palette.items() if v==color),"palette":palette,"binding":"fill"}
        fixed=[11,13,15,17]
    elif key=="heatmap_table":
        palette={"高":"5A3921","中":"C5A681","低":"EFEEE8","なし":"FFFFFF"}
        for sid in [20,22,24,26,29,31,33,35,38,40,42,44,47,49,51,53]:
            states[f"state_{sid:03}"]={"shape":sm[sid].name,"label":sm[sid+1].name,"sample":sm[sid+1].text,"palette":palette,"binding":"fill"}
        fixed=[56,58,60,62]
    elif key=="harvey_ball_table":
        # Each marker has one invariant circle and one invariant half-circle.
        # Fill changes preserve its cell position and bound.
        groups=[([18],"満たす"),([19,20],"一部満たす"),([21,22],"一部満たす"),([23],"満たさない"),([24,25],"一部満たす"),([26],"満たす")]
        pie_proto=copy.deepcopy(sm[19].element); circle_proto=copy.deepcopy(sm[23].element)
        max_id=max(sm)
        for i,(ids,state) in enumerate(groups):
            reference=sm[ids[0]]; x,y,w,h=box(reference)
            pair=[]
            for proto in [circle_proto,pie_proto]:
                element=copy.deepcopy(proto); max_id+=1
                element.xpath('.//p:cNvPr')[0].set('id',str(max_id)); element.xpath('.//p:cNvPr')[0].set('name',f'shape:{max_id:03}')
                off=element.xpath('.//a:xfrm/a:off')[0]; off.set('x',str(x));off.set('y',str(y))
                ext=element.xpath('.//a:xfrm/a:ext')[0];ext.set('cx',str(w));ext.set('cy',str(h))
                slide.shapes._spTree.insert_element_before(element,'p:extLst');pair.append(f'shape:{max_id:03}')
            remove_ids(slide,ids)
            states[f"rating_{i+1}"]={"shape":pair[0],"half_shape":pair[1],"sample":state,
                                     "palette":{"満たす":"322014","一部満たす":"322014","満たさない":"FFFFFF"},"binding":"harvey"}
        fixed=[11,14,16]
    for sid in fixed: sm[sid].name=f"static:legend_{sid}"
    return states


def build():
    inventory = json.loads((ROOT / "catalog/inventory.json").read_text(encoding="utf-8"))
    upstream = Presentation(str(ROOT / "catalog/upstream/assets/SuperTemplate_62type.pptx"))
    mapping = {}
    for index, slide in enumerate(upstream.slides, 1):
        text = "\n".join(s.text for s in slide.shapes if s.has_text_frame)
        match = re.search(r"パーツ(\d+)｜",text)
        if match: mapping[f"b{int(match[1]):02}"] = index
        else:
            for entry in inventory["layouts"]:
                if re.search(r"\d\d  " + re.escape(entry["layout_id"].upper()) + r"\b", text):
                    mapping[entry["part_id"]] = index
    assert len(mapping) == 62, mapping
    entries = []
    for record in inventory["layouts"]:
        entry = copy.deepcopy(record)
        index = mapping[entry["part_id"]]
        p = Presentation(); p.slide_width=upstream.slide_width; p.slide_height=upstream.slide_height
        s = clone_slide(p,upstream.slides[index-1])
        comparison=json.loads((ROOT/"catalog/html-comparison.json").read_text(encoding="utf-8")) if (ROOT/"catalog/html-comparison.json").exists() else {}
        changes = ["Catalog header/footer are generated metadata; content remains in native shapes at the source positions.",
                   "Source-preserved compact type sizes are a separate strict catalog theme; legacy 8 layouts retain their 22–28pt body policy."]
        changes.extend(comparison.get(entry["layout_id"],[]))
        chart_upgrade(s,entry["layout_id"],changes)
        if entry["layout_id"]=="calc_flow":
            for sh in s.shapes:
                if sh.shape_id in (22,34):
                    sh.left-=Inches(.125); sh.width+=Inches(.25)
            changes.append("Widened operator boxes by 0.25in within the existing gap after PowerPoint measured the multiplication glyph overflow. Calculation data now enforces A×B=base+increment.")
        repaired=0
        for paragraph in s._element.xpath(".//a:p"):
            for redundant in paragraph.findall(qn("a:pPr"))[1:]:
                paragraph.remove(redundant); repaired+=1
        if repaired: changes.append(f"Repaired {repaired} redundant paragraph-property nodes from source OOXML; one a:pPr per paragraph.")
        for shape in s.shapes:
            shape.name = f"shape:{shape.shape_id:03}"
            if shape.has_text_frame:
                shape.text_frame.auto_size = MSO_AUTO_SIZE.NONE
                # Keep the source's line breaks and styles; no autoshrink.
            if shape.has_chart:
                for node in shape.chart._chartSpace.xpath(".//a:spAutoFit|.//a:normAutofit"):
                    node.getparent().remove(node)
        entry["metrics"]={}
        configure_metrics(s,entry)
        entry["states"]=configure_states(s,entry)
        metric_labels = {v.get("label") for v in entry["metrics"].values()}
        state_labels = {v.get("label") for v in entry["states"].values()}
        candidates = [x for x in s.shapes if x.has_text_frame and x.text and 0.8 < x.top/Inches(1) < 6.8]
        title = max(candidates,key=lambda x:max([r.font.size.pt for pp in x.text_frame.paragraphs for r in pp.runs if r.font.size] or [0]))
        # Regular title is always shape 5; big numeric figures must not steal it.
        if any(x.shape_id==5 and abs(x.top/Inches(1)-1)<.1 for x in candidates):
            title = next(x for x in candidates if x.shape_id==5)
        title.name="slot:title"
        entry["title"]= {"shape":title.name,"sample":title.text,"bounds_emu":box(title),**frame_spec(title.text_frame,title.width,title.height)}
        texts, charts = {}, {}
        for shape in s.shapes:
            if shape.has_chart:
                charts[f"chart_{shape.shape_id:03}"] = chart_spec(shape)
            if shape.has_table:
                for row_index,row in enumerate(shape.table.rows):
                    for col_index,cell in enumerate(row.cells):
                        if cell.is_spanned or not cell.text: continue
                        name=f"table_{shape.shape_id:03}_r{row_index}_c{col_index}"
                        width=sum(shape.table.columns[j].width for j in range(col_index,col_index+cell.span_width))
                        height=sum(shape.table.rows[j].height for j in range(row_index,row_index+cell.span_height))
                        texts[name]={"shape":shape.name,"cell":[row_index,col_index],"sample":cell.text,
                                     **frame_spec(cell.text_frame,width,height)}
            if not shape.has_text_frame or not shape.text or shape==title or shape.name in metric_labels|state_labels or shape.name.startswith("static:"): continue
            y=shape.top/Inches(1)
            if y < .7:
                shape.name="meta:header" if shape.left/Inches(1)<1 else "meta:layout"
            elif y > 7.0:
                shape.name="meta:page" if shape.left/Inches(1)>12 else "meta:format" if shape.left/Inches(1)>10 else "meta:footer"
            elif entry["layout_id"]=="calc_flow" and shape.shape_id in (22,34):
                shape.name=f"static:operator_{shape.shape_id}"
            else:
                texts[f"text_{shape.shape_id:03}"]={"shape":shape.name,"sample":shape.text,"bounds_emu":box(shape),
                                                  **frame_spec(shape.text_frame,shape.width,shape.height)}
        entry["charts"]=charts
        entry["chart_aliases"]=configure_aliases(s,entry)
        texts={k:v for k,v in texts.items() if v["shape"] not in entry["chart_aliases"]}
        if entry["layout_id"]=="scenario_lines_cagr":texts["text_012"]["sample"]="単位、2026〜2031年"
        if entry["chart_aliases"]:
            changes.append("Chart category/series labels and aggregate callouts are bound to data. Totals/CAGR recompute; source example inconsistencies are corrected.")
        entry.update(texts=texts,charts=charts,source_pptx_slide=index,template=f"catalog/templates/{entry['layout_id']}.pptx",
                     status="template_created",differences=changes,split_policy="reject_overflow; explicit replan with retained origin and references",
                     image_policy="no image slots in this native source catalog; image input rejected; use legacy text_image for fit/crop",
                     editability={"text":"native", "tables":"native cells where table exists; remaining matrices are native shapes",
                                  "charts":"native chart with embedded workbook where registered; metric-bound diagrams use native shapes",
                                  "shapes":"native", "pictures":0},
                     unsupported="No arbitrary geometry, image insertion, automatic semantic pagination or unsupported chart types.")
        p.core_properties.author="slide-layout-agent / Carnot AI Inc."
        p.core_properties.title=entry["name"]
        p.core_properties.comments="Derived from Carnot AI Inc. MIT-licensed SuperTemplate; see catalog/upstream/LICENSE."
        s.notes_slide.notes_text_frame.text=json.dumps({"layout_id":entry["layout_id"],"upstream_commit":inventory["commit"],
                                                      "source_pptx_slide":index,"license":(ROOT/"catalog/upstream/LICENSE").read_text(encoding="utf-8")})
        dest=ROOT / entry["template"]; dest.parent.mkdir(parents=True,exist_ok=True); p.save(dest)
        entry["template_sha256"]=hashlib.sha256(dest.read_bytes()).hexdigest()
        entry["shape_count"]=len(s.shapes)
        entry["native_table_count"]=sum(x.has_table for x in s.shapes)
        entry["native_chart_count"]=sum(x.has_chart for x in s.shapes)
        entry["capacity"]={"text_slots":len(texts),"chart_slots":len(charts),"metric_slots":len(entry["metrics"]),
                            "state_slots":len(entry["states"]),
                            "required":"exact slot keys; fixed table dimensions/chart cardinality; source font sizes; no shrink"}
        entries.append(entry)
    manifest={"version":1,"upstream_commit":inventory["commit"],"layout_count":len(entries),"variants":["warm","cool"],
              "license_notice":"catalog/upstream/LICENSE","layouts":entries}
    write(ROOT / "catalog/manifest.json",manifest)
    from slide_agent.schema_contract import layout_schema
    from slide_agent.models import Plan
    for entry in entries:write(ROOT/f"catalog/schemas/{entry['layout_id']}.schema.json",layout_schema(entry))
    write(ROOT/'catalog/plan.schema.json',Plan.model_json_schema())
    print(f"Built {len(entries)} templates, {sum(e['native_chart_count'] for e in entries)} native charts")


if __name__ == "__main__": build()
