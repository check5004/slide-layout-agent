"""Build original native PPTX templates for everyday research sharing.

Generation consumes these saved files; it never runs this geometry builder.
All samples are fictional. No Office or network is used here.
"""
import argparse
import hashlib
import json
import math
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_DATA_LABEL_POSITION, XL_TICK_MARK
from pptx.enum.text import MSO_AUTO_SIZE, MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from pptx.oxml.xmlchemy import OxmlElement
from slide_agent.editorial_specs import SPECS

OUT = ROOT / 'catalog/editorial'
INK, MUTED, TEAL, BLUE = '213547', '566777', '087E83', '4867A8'
BG, SURFACE, LINE, TINT = 'F7F8FA', 'FFFFFF', 'DCE3E9', 'E8F3F2'
FONT = 'Yu Gothic'
SAMPLE_HEADS = ['検索入口を揃える', '根拠まで辿れる', '更新範囲を決める', '例外を見える化', '比較軸を共有する', '確認先を残す']
SAMPLE_BODIES = [
    '用語の揺れを減らし、部門をまたいでも同じ資料へ辿れる導線を整える。検索結果には対象範囲を添え、読み手が必要な情報かを判断できるようにする。',
    '要約だけで判断せず、原文の条件と確認日を一緒に読み取れる形にする。引用先の位置も残し、解釈に疑問があれば読み手が前提を確かめられるようにする。',
    '変更頻度に応じて確認担当を置き、古い記述を見分けられるようにする。更新が止まった資料は利用条件を明記し、再確認の対象として管理する。',
    '適用できない場面も記録し、利用者が前提の違いを確認できるようにする。例外を整理したうえで、運用上の判断が必要な範囲を担当者と共有する。',
    '評価項目と対象範囲を先に揃え、異なる条件の結果を混同しない。資料間で一致しない定義を残し、単純な優劣として扱えるかを確認する。',
    '未確認の論点を残し、追加調査で何を確かめるかを明確にする。次の担当者が判断の経緯を追えるよう、保留した理由と確認先も記録する。',
]
SHORT_BODIES = ['用語を揃え、同じ資料へ辿れる入口を整える。',
                '原文の条件と確認日を、要約とともに残す。',
                '変更頻度に応じて担当と再確認の時期を決める。',
                '適用できない場面も記録し、前提の違いを残す。',
                '評価項目を揃え、異なる条件の結果を混同しない。',
                '未確認の論点と、追加調査で確かめる点を残す。']


def write(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.is_file() and json.loads(path.read_text(encoding='utf-8')) == value:
        return
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


def style(tf, text, size, color=INK, bold=False):
    tf.clear(); tf.word_wrap = True; tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = Inches(.02)
    tf.margin_top = tf.margin_bottom = Inches(.015)
    for index, line in enumerate(text.split('\n')):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        for key,value in [('hangingPunct','0'),('eaLnBrk','1'),('latinLnBrk','0')]:
            p._p.get_or_add_pPr().set(key,value)
        # Explicit point leading is stable for Japanese fonts; percentage
        # leading otherwise depends on the viewer's CJK font metrics.
        p.line_spacing = Pt(size * 1.25); p.space_before = p.space_after = Pt(0)
        run = p.add_run(); run.text = line
        for font in (run.font, p.font):
            font._rPr.set('lang', 'ja-JP')
            font.name = FONT; font.size = Pt(size); font.bold = bold
            font.color.rgb = RGBColor.from_string(color)
            for tag in ('a:ea', 'a:cs'):
                node = font._rPr.find(tag, font._rPr.nsmap)
                if node is None:
                    node = OxmlElement(tag); font._rPr.append(node)
                node.set('typeface', FONT)


def capacity(w, h, size):
    # Explicit conservative bound, with room for Japanese punctuation/wrapping.
    return dict(font_pt=size, font_family=FONT,
                max_units_per_line=round(max(1, (w * 72 - 6) / (size * 1.08)), 2),
                max_lines=max(1, 1 + math.floor((h * 72 - 2.2 - size * 1.33) / (size * 1.25))),
                overflow='reject; split or replan; never shrink')


class Builder:
    def __init__(self, spec):
        self.entry = {**spec, 'catalog': 'editorial', 'part_id': spec['layout_id'],
                      'source_file': 'original fictional research-sharing design', 'source_section': 1,
                      'texts': {}, 'charts': {}, 'metrics': {}, 'states': {}, 'images': {}, 'text_flow': [], 'chart_aliases': {}, 'rows': {}}
        self.prs = Presentation(); self.prs.slide_width = Inches(13.333333); self.prs.slide_height = Inches(7.5)
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.slide.background.fill.solid(); self.slide.background.fill.fore_color.rgb = RGBColor.from_string(BG)
        self.counter = 0
        self.rect(.62, .43, .32, .055, TEAL)
        self.text('meta:header', '架空調査｜社内ナレッジ共有', 1.05, .32, 9.9, .32, 12, MUTED)
        self.text('meta:page', '01', 12.0, 6.99, .68, .3, 12, MUTED)
        self.text('meta:footer', 'FICTIONAL RESEARCH / EDITORIAL', .66, 7.02, 10.5, .26, 10, MUTED)
        self.rect(.66, 6.85, 12.0, .012, LINE)

    def rect(self, x, y, w, h, color, border=None):
        self.counter += 1
        shape = self.slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.name = f'decor:{self.counter}'
        shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string(color)
        if border:
            shape.line.color.rgb = RGBColor.from_string(border); shape.line.width = Pt(.7)
        else: shape.line.fill.background()
        shape._element.spPr.append(OxmlElement('a:effectLst'))
        return shape

    def text(self, name, text, x, y, w, h, size=20, color=INK, bold=False):
        shape = self.slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        shape.name = name; style(shape.text_frame, text, size, color, bold)
        return shape

    def slot(self, key, text, x, y, w, h, size=20, color=INK, bold=False, role='body'):
        shape = self.text('slot:' + key, text, x, y, w, h, size, color, bold)
        spec = {'shape': shape.name, 'sample': text, 'role': role, 'vertical_anchor':'top',
                'bounds_emu': [shape.left, shape.top, shape.width, shape.height], **capacity(w, h, size)}
        if key == 'title': self.entry['title'] = spec
        else: self.entry['texts'][key] = spec
        return shape

    def row_group(self, key, keys, y, h, static_names=()):
        """Center parallel fields; keep body/evidence as one top-aligned stack.

        Text capacity remains the previously authored bound, even when a short
        label receives the full row-height frame for native middle anchoring.
        """
        top, height = Inches(y + .03), Inches(h - .06)
        row = self.entry['rows'].setdefault(key, {'top_emu':top,'height_emu':height,
            'default_alignment':'middle','members':[]})
        flows = [f for f in self.entry['text_flow'] if f['body'] in keys and f['following'] in keys]
        flow_keys = {f[k] for f in flows for k in ('body','following')}
        for flow in flows: flow['row'] = key
        shapes = {s.name:s for s in self.slide.shapes}
        for slot_key in keys:
            spec = self.entry['texts'][slot_key];shape = shapes[spec['shape']]
            row['members'].append({'shape':shape.name})
            if slot_key not in flow_keys:
                shape.top, shape.height = top, height
                shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
                spec['vertical_anchor'] = 'middle'
                spec['bounds_emu'] = [shape.left,shape.top,shape.width,shape.height]
        for name in static_names:
            shape=shapes[name];shape.top,shape.height=top,height
            shape.text_frame.vertical_anchor=MSO_ANCHOR.MIDDLE
            row['members'].append({'shape':name})

    def header(self, title=None):
        self.slot('title', title or '社内ナレッジを探しやすくする条件', .66, .86, 12.0, .66, 28, bold=True, role='title')
        self.slot('context', '架空の社内調査。運用条件を整理した所見であり、効果を実証したものではない。',
                  .68, 1.63, 11.95, .55, 15, MUTED, role='context')

    def card(self, i, x, y, w, h, featured=False, compact=False, evidence=True):
        fill = TINT if featured else SURFACE
        self.rect(x, y, w, h, fill, LINE)
        self.rect(x, y, .045, h, TEAL if featured else LINE)
        pad = .23
        small = h < 2.5
        self.slot(f'item_{i}_heading', SAMPLE_HEADS[i - 1], x + pad, y + .15, w - 2 * pad, .48 if small else .7,
                  21 if featured else 20, TEAL if featured else INK, True, 'heading')
        body_y = y + (.72 if small else 1.02)
        foot_h = .49 if evidence else 0
        body_h = h - (body_y - y) - foot_h - .10
        if featured: body_h = min(body_h, 1.8)
        self.slot(f'item_{i}_body', SHORT_BODIES[i - 1] if compact else SAMPLE_BODIES[i - 1],
                  x + pad, body_y, w - 2 * pad, body_h, 18 if compact else 20)
        if evidence:
            evidence_y = body_y + body_h + .12
            evidence = ['対象：資料の入口を確認', '根拠：原文の条件を保持', '留保：運用負荷は未検証',
                        '注意：対象外の場面を含む', '根拠：比較条件を照合', '保留：担当部署へ確認'][i-1]
            self.slot(f'item_{i}_evidence', evidence, x + pad, evidence_y,
                      w - 2 * pad, .34, 14, MUTED, role='evidence')
            self.entry['text_flow'].append({'body': f'item_{i}_body', 'following': f'item_{i}_evidence', 'gap_emu': Inches(.12)})

    def row(self, i, x, y, w, h, evidence=False, numbered=False):
        self.rect(x, y + h - .02, w, .012, LINE)
        if w < 8 and not evidence:
            self.slot(f'item_{i}_heading', SAMPLE_HEADS[i-1], x + .06, y + .02, w - .12, .45, 20, INK, True, 'heading')
            self.slot(f'item_{i}_body', SHORT_BODIES[i-1], x + .06, y + .57, w - .12, h - .64, 18)
            return
        if w < 8 and evidence:
            dense = h < 1.55
            self.slot(f'item_{i}_heading', SAMPLE_HEADS[i-1], x + .06, y + .02, w - .12, .46, 20, INK, True, 'heading')
            self.slot(f'item_{i}_body', SHORT_BODIES[i-1] if not dense else ['同じ資料へ辿れる入口を整える。', '原文の条件と確認日を残す。', '更新担当と確認時期を決める。', '対象外の場面と前提を記録する。'][i-1],
                      x + .06, y + (.50 if dense else .58), w - .12, h - (.91 if dense else 1.0), 18)
            self.slot(f'item_{i}_evidence', ['対象：検索導線', '根拠：原文の条件', '留保：運用負荷は未検証', '注意：適用範囲を確認'][i-1],
                      x + .06, y + h - .35, w - .12, .30, 13, MUTED, role='evidence')
            self.entry['text_flow'].append({'body':f'item_{i}_body','following':f'item_{i}_evidence','gap_emu':Inches(.08)})
            return
        offset = .57 if numbered else 0
        if numbered:
            self.text(f'static:number_{i}', f'{i:02}', x, y + .03, .45, .4, 17, TEAL, True)
        # Heading and prose form columns within each record, not two unrelated lists.
        side_note = evidence and h < 1.05
        heading_w = min(4.05 if side_note else 3.35 if numbered else 3.10, w * .34)
        heading = ['検索の入口を部署間で揃える','原文の根拠と条件へ辿れる','更新範囲と担当を明確にする','適用できない場面も記録する','比較する条件を共有する','次に確認する論点を残す'][i-1] if side_note else SAMPLE_HEADS[i - 1]
        self.slot(f'item_{i}_heading', heading, x + offset, y + .02,
                  heading_w - offset, min(.75, h - .03), 18 if side_note else 20, INK, True, 'heading')
        body = ['用語を揃え、\n同じ資料へ辿れる入口を整える。', '原文の条件と確認日を、\n要約とともに残す。',
                '変更頻度に応じて、\n担当と再確認の時期を決める。', '適用できない場面も記録し、\n前提の違いを残す。',
                '評価項目を揃え、\n異なる条件の結果を混同しない。', '未確認の論点と、\n追加調査で確かめる点を残す。'][i-1] if side_note else SHORT_BODIES[i-1]
        self.slot(f'item_{i}_body', body, x + heading_w + .23, y + .02,
                  w - heading_w - .23 - (2.2 if side_note else 0), h - (.47 if evidence and not side_note else .03), 18)
        if evidence:
            self.slot(f'item_{i}_evidence', f'根拠：架空資料{chr(64+i)}\n確認：日付未記録' if side_note else '留保：追加の確認が必要', x + w - 2.05 if side_note else x + heading_w + .23,
                      y + .03 if side_note else y + h - .39, 2.0 if side_note else w - heading_w - .23,
                      h - .07 if side_note else .32, 12, MUTED, role='evidence')
            if not side_note:
                self.entry['text_flow'].append({'body':f'item_{i}_body','following':f'item_{i}_evidence','gap_emu':Inches(.10)})
        self.row_group(f'item_{i}', [f'item_{i}_heading',f'item_{i}_body'] +
                       ([f'item_{i}_evidence'] if evidence else []), y, h,
                       [f'static:number_{i}'] if numbered else [])


def build_one(spec):
    b = Builder(spec); f, v, n = spec['family'], spec['structural_variant'], spec['item_count']
    if f == 'cover':
        if v == 'brief':
            b.slot('title', '社内ナレッジ共有の調査メモ', .7, 1.2, 10.8, 1.25, 34, bold=True, role='title')
            b.slot('context', '検索性と信頼性を両立するための運用条件', .72, 2.65, 11.6, .75, 23, MUTED, role='context')
            for i in range(1, 3): b.card(i, .7 + (i - 1) * 6.1, 4.0, 5.92, 2.4, compact=True, evidence=False)
        elif v == 'split':
            b.rect(8.3, .85, 4.38, 5.72, TINT)
            b.slot('title', '社内ナレッジ共有の調査メモ', .7, 1.5, 7.05, 1.65, 34, bold=True, role='title')
            b.slot('context', '検索性と信頼性を両立するための運用条件', .72, 3.4, 6.85, 1.2, 23, MUTED, role='context')
            for i in range(1, 3): b.card(i, 8.53, 1.12 + (i - 1) * 2.63, 3.9, 2.43, compact=True, evidence=False)
        else:
            b.rect(.66, 1.0, 12.0, 2.6, INK)
            b.slot('title', '社内ナレッジ共有の調査メモ', .94, 1.3, 11.4, 1.1, 34, 'FFFFFF', True, 'title')
            b.slot('context', '検索性と信頼性を両立するための運用条件', .96, 2.6, 11.35, .65, 23, 'FFFFFF', role='context')
            for i in range(1, 3): b.row(i, .8, 4.03 + (i - 1) * 1.17, 11.6, 1.05)
    else:
        b.header()
        if f == 'cards':
            if v.startswith('image_'):
                ix, tx = (.66, 6.86) if v == 'image_left' else (6.86, .66)
                frame = b.rect(ix, 2.34, 5.8, 4.22, 'E3EAF0'); frame.name = 'slot:image'
                b.entry['images']['image'] = {'shape': frame.name, 'bounds_emu': [frame.left, frame.top, frame.width, frame.height],
                                             'modes': ['fit', 'crop'], 'required': True}
                for k in range(2): b.card(k + 1, tx, 2.34 + k * 2.21, 5.8, 2.01, compact=True)
            elif v == 'grid':
                cols = n if n <= 3 else 2 if n == 4 else 3
                rows = math.ceil(n / cols); w = (12 - .22 * (cols - 1)) / cols; h = (4.22 - .22 * (rows - 1)) / rows
                for k in range(n): b.card(k + 1, .66 + (k % cols) * (w + .22), 2.34 + (k // cols) * (h + .22), w, h, compact=rows > 1)
            elif v == 'rows':
                h = 4.28 / n
                for k in range(n): b.row(k + 1, .75, 2.31 + k * h, 11.85, h, evidence=True)
            else:
                featured_x, others_x = (.66, 6.85) if v == 'feature_left' else (6.85, .66)
                b.card(1, featured_x, 2.34, 5.81, 4.22, featured=True)
                h = (4.22 - .17 * (n - 2)) / (n - 1)
                for k in range(n - 1): b.row(k + 2, others_x, 2.36 + k * (h + .17), 5.81, h, evidence=True)
        elif f in ('bullets', 'agenda', 'summary'):
            if v in ('rows', 'rail'):
                x, w = (2.0, 10.6) if v == 'rail' else (.76, 11.8)
                if v == 'rail': b.rect(.68, 2.32, 1.0, 4.25, TINT)
                for k in range(n): b.row(k + 1, x, 2.36 + k * 4.22 / n, w, 4.22 / n, numbered=f == 'agenda')
            elif v == 'split':
                b.card(1, .66, 2.34, 5.8, 4.22, featured=True, evidence=False)
                for k in range(2): b.card(k + 2, 6.75, 2.34 + k * 2.21, 5.91, 2.02, compact=True, evidence=False)
            else:
                cols = min(n, 3); rows = math.ceil(n / cols)
                w, h = (12 - .3 * (cols - 1)) / cols, (4.22 - .24 * (rows - 1)) / rows
                for k in range(n):
                    x, y = .66 + (k % cols) * (w + .3), 2.34 + (k // cols) * (h + .24)
                    b.rect(x, y, w, .045, TEAL)
                    b.slot(f'item_{k+1}_heading', SAMPLE_HEADS[k], x, y + .2, w, .72, 21, bold=True, role='heading')
                    b.slot(f'item_{k+1}_body', SHORT_BODIES[k], x, y + 1.03, w, h - 1.12, 18 if rows > 1 else 20)
        elif f == 'takeaway':
            if v in ('left', 'right'):
                x, rx = (.66, 6.22) if v == 'left' else (7.5, .66)
                b.rect(x, 2.34, 5.16, 4.22, TINT)
                b.slot('conclusion', '検索の入口だけでなく、根拠と更新責任を一緒に設計する', x + .27, 2.8, 4.62, 2.7, 27, TEAL, True)
                for k in range(3): b.row(k + 1, rx, 2.39 + k * 1.4, 6.42, 1.32)
            else:
                h = 1.38 if v == 'top' else .95
                b.rect(.66, 2.33, 12, h, TINT)
                b.slot('conclusion', '検索の入口だけでなく、根拠と更新責任を一緒に設計する', .92, 2.51, 11.45, h - .24, 25, TEAL, True)
                y = 2.33 + h + .25
                for k in range(3): b.card(k + 1, .66 + k * 4.08, y, 3.84, 6.56 - y, compact=True, evidence=False)
        elif f == 'diagram':
            wide = v.startswith('wide'); dw = 7.15 if wide else 5.7
            dx, tx = (.66, .66 + dw + .42) if v.endswith('left') else (12.66 - dw, .66)
            tw = 12 - dw - .42
            b.rect(dx, 2.34, dw, 4.22, TINT)
            for k in range(3):
                y = 2.57 + k * 1.26
                b.rect(dx + .27, y, dw - .54, .96, SURFACE)
                b.slot(f'node_{k+1}', ['入力資料', '条件の確認', '所見の共有'][k], dx + .5, y + .08, dw - 1.0, .42, 21, TEAL, True, 'node')
                b.slot(f'node_{k+1}_detail', ['原文・確認日・対象範囲を受け取る', '定義と比較条件の違いを整理する', '根拠と未確認の論点を残して渡す'][k], dx + .5, y + .57, dw - 1.0, .31, 14, MUTED, role='node_detail')
                if k < 2: b.text(f'static:arrow_{k}', '↓', dx + dw / 2 - .2, y + .96, .5, .3, 15, TEAL)
            b.slot('explanation_heading', '原文とのつながりを残す', tx, 2.54, tw, .84, 23, bold=True, role='heading')
            b.slot('explanation', '資料を集めた後で、対象範囲と比較条件を確認する。共有時には原文の位置を添え、読み手が結論の前提へ戻れるようにする。', tx, 3.57, tw, 2.15, 20)
            b.slot('caveat', '留保：因果関係を示す工程図ではない。', tx, 6.0, tw, .52, 14, MUTED, role='evidence')
            b.entry['text_flow'].append({'body':'explanation','following':'caveat','gap_emu':Inches(.16)})
        elif f == 'comparison':
            points = {'a': ['入口を統一し、部署をまたぐ検索をしやすくする。', '更新作業が担当者に集中するため、確認待ちの滞留に注意する。'],
                      'b': ['各部署が原文を管理し、専門用語や例外の説明を保ちやすい。', '横断検索には共通の索引が必要。部署ごとの分類の違いを整理する。']}
            # A/B identity follows its labels; swapping never reverses chronological data.
            if v in ('columns', 'swapped'):
                for k, key in enumerate(('a', 'b') if v == 'columns' else ('b', 'a')):
                    x = .66 + k * 6.15; b.rect(x, 2.34, 5.85, 4.22, SURFACE, LINE)
                    b.slot(f'{key}_heading', '集約型' if key == 'a' else '分散型', x + .25, 2.57, 5.35, .62, 23, TEAL, True, 'heading')
                    for j in range(2):
                        b.slot(f'{key}_axis_{j+1}', ['検索上の特徴','運用上の注意'][j], x+.25, 3.28+j*1.56, 5.35, .32, 14, MUTED, True, 'axis')
                        b.slot(f'{key}_point_{j+1}', points[key][j], x + .25, 3.68 + j * 1.56, 5.35, 1.12, 20)
            else:
                for k, key in enumerate(('a', 'b')):
                    y = 2.34 + k * 2.22; b.rect(.66, y, 12, 2.0, SURFACE, LINE)
                    b.slot(f'{key}_heading', '集約型' if key == 'a' else '分散型', .91, y + (.15 if v=='matrix' else .04), 2.4 if v=='matrix' else 11.3, .5, 23, TEAL, True, 'heading')
                    for j in range(2):
                        x = 3.6 + j * 4.4 if v=='matrix' else .92+j*5.95
                        b.slot(f'{key}_axis_{j+1}', ['検索上の特徴','運用上の注意'][j], x, y+(.14 if v=='matrix' else .55), 4.02 if v=='matrix' else 5.55, .28, 14, MUTED, True, 'axis')
                        b.slot(f'{key}_point_{j+1}', points[key][j], x, y + (.53 if v=='matrix' else .86), 4.02 if v=='matrix' else 5.55, 1.37 if v=='matrix' else 1.11, 20)
                if v == 'matrix':
                    # Same semantic cells, tighter headings and a visible column boundary.
                    b.rect(7.74, 2.34, .025, 4.22, LINE)
                    b.rect(3.34, 2.34, .025, 4.22, LINE)
            b.entry['comparison_axes'] = [['a_axis_1','b_axis_1'],['a_axis_2','b_axis_2']]
        elif f in ('steps', 'timeline'):
            if v == 'horizontal':
                w = (12 - .34 * (n - 1)) / n
                for k in range(n):
                    x = .66 + k * (w + .34)
                    b.text(f'static:sequence_{k}', f'{k+1:02}', x, 2.46, w, .39, 17, TEAL, True)
                    b.rect(x, 2.99, w, .04, LINE)
                    b.slot(f'item_{k+1}_heading', SAMPLE_HEADS[k], x, 3.18, w, .85, 21, bold=True, role='heading')
                    b.slot(f'item_{k+1}_body', SHORT_BODIES[k], x, 4.09, w, 2.20, 19)
                    if f == 'timeline': b.slot(f'item_{k+1}_date', ['準備期', '試行期', '評価期', '展開期'][k], x, 2.35, w, .7, 19, TEAL, True, 'date')
                    if f == 'timeline':
                        sh = next(s for s in b.slide.shapes if s.name == f'static:sequence_{k}'); sh.text_frame.clear()
            else:
                b.rect(2.10 if f == 'timeline' else .94, 2.51, .025, 3.9, LINE)
                h = 4.25 / n
                for k in range(n):
                    y = 2.34 + k * h
                    b.row(k + 1, 2.42 if f == 'timeline' else 1.35, y, 10.18 if f == 'timeline' else 11.25, h, numbered=f == 'steps')
                    if f == 'timeline':
                        b.slot(f'item_{k+1}_date', ['準備期', '試行期', '評価期', '展開期'][k], .68, y, 1.28, .65, 16, TEAL, True, 'date')
                        b.row_group(f'item_{k+1}',[f'item_{k+1}_date'],y,h)
        elif f == 'metrics':
            if v.startswith('breakdown'):
                top = v.endswith('top')
                tx,ty,tw,th=(.66,2.34,12,1.1) if top else (.66,2.34,3.12,4.22)
                b.rect(tx,ty,tw,th,TINT)
                b.slot('total_heading','対象資料の総数',tx+.23,ty+.18,3.05 if top else 2.66,.54,20,TEAL,True,'heading')
                total=b.slot('total_label','120件',tx+3.4 if top else tx+.23,ty+.12 if top else ty+1.03,2.7,.84,36,TEAL,True)
                cap=b.entry['texts'].pop('total_label')
                b.entry['metrics']['total']=dict(binding='label',label=total.name,suffix='件',min=0,max=1e12,sample=120,label_capacity=cap)
                b.slot('total_note','各資料を一つの状態に分類した架空の棚卸し。',tx+6.65 if top else tx+.23,ty+.18 if top else ty+2.25,4.95 if top else 2.66,.69 if top else 1.52,16,MUTED,role='evidence')
                for k in range(3):
                    x,y,w,h=(.66+k*4.08,3.73,3.84,2.83) if top else (4.08,2.34+k*1.44,8.58,1.30)
                    b.rect(x,y,w,h,SURFACE,LINE)
                    b.slot(f'item_{k+1}_heading',['確認済み','要確認','更新予定'][k],x+.22,y+.16, w-.44 if top else 2.1,.49,20,bold=True,role='heading')
                    num=b.slot(f'metric_label_{k+1}',str([84,24,12][k])+'件',x+.22 if top else x+2.4,y+.89 if top else y+.2,2.2,.71,32,TEAL,True)
                    b.slot(f'item_{k+1}_body',['対象範囲と担当を確認済み。','条件か確認日の照合が必要。','改訂の予定を登録済み。'][k],x+.22 if top else x+4.85,y+1.93 if top else y+.20,w-.44 if top else 3.5,.7 if top else .90,18)
                    if not top: b.row_group(f'item_{k+1}',[f'item_{k+1}_heading',f'metric_label_{k+1}',f'item_{k+1}_body'],y,h)
                    cap=b.entry['texts'].pop(f'metric_label_{k+1}')
                    b.entry['metrics'][f'metric_{k+1}']=dict(binding='label',label=num.name,suffix='件',min=0,max=1e12,sample=[84,24,12][k],label_capacity=cap)
                b.entry['metric_total']={'total':'total','parts':['metric_1','metric_2','metric_3']}
                b.entry['texts']['context']['sample']='架空の資料一覧を、相互に重ならない確認状態で整理した例。業務への効果は未検証。'
                style(next(s for s in b.slide.shapes if s.name=='slot:context').text_frame,b.entry['texts']['context']['sample'],15,MUTED)
            for k in range(n):
                if v.startswith('breakdown'): break
                x, y, w, h = (.66 + k * 12.24 / n, 2.34, 12.24 / n - .24, 4.22) if v == 'row' else (.66, 2.34 + k * 1.44, 12, 1.3)
                b.rect(x, y, w, h, SURFACE, LINE)
                if v == 'row':
                    b.slot(f'item_{k+1}_heading', ['対象範囲の記載', '確認担当の記載', '更新日の記載', '留保条件の記載'][k], x + .22, y + .25, w - .44, .7, 20, bold=True, role='heading')
                    num = b.slot(f'metric_label_{k+1}', str([120, 84, 24, 12][k])+'件', x + .22, y + 1.25, w - .44, .85, 38, TEAL, True)
                    b.slot(f'item_{k+1}_body', ['適用できる業務範囲が記載された文書。', '内容を確認する担当が記載された文書。', '原文の最終確認日が記載された文書。', '未検証の条件や対象外の場面を記載した文書。'][k], x + .22, y + 2.4, w - .44, 1.18, 18)
                else:
                    b.slot(f'item_{k+1}_heading', ['対象範囲の記載', '確認担当の記載', '更新日の記載'][k], x + .25, y + .27, 2.8, .7, 20, bold=True, role='heading')
                    num = b.slot(f'metric_label_{k+1}', str([120, 84, 24][k])+'件', x + 3.3, y + .18, 2.0, .87, 38, TEAL, True)
                    b.slot(f'item_{k+1}_body', ['適用できる業務範囲が記載された文書。', '内容を確認する担当が記載された文書。', '原文の最終確認日が記載された文書。'][k], x + 5.7, y + .2, 5.9, .92, 18)
                    b.row_group(f'item_{k+1}',[f'item_{k+1}_heading',f'metric_label_{k+1}',f'item_{k+1}_body'],y,h)
                cap = b.entry['texts'].pop(f'metric_label_{k+1}')
                b.entry['metrics'][f'metric_{k+1}'] = dict(binding='label', label=num.name, suffix='件', min=0, max=1e12, sample=[120, 84, 24, 12][k], label_capacity=cap)
            if not v.startswith('breakdown'):
                b.entry['texts']['context']['sample'] = '架空集計：資料120件の記載状況。項目間には重複があり、合計を総数として扱わない。'
                style(next(s for s in b.slide.shapes if s.name=='slot:context').text_frame, b.entry['texts']['context']['sample'], 15, MUTED)
        elif f == 'table':
            cols = 3 if v == 'three_columns' else 2
            shape = b.slide.shapes.add_table(n + 1, cols, Inches(.66), Inches(2.35), Inches(12), Inches(3.9)); shape.name = 'slot:table'
            for r in range(n + 1):
                for c in range(cols):
                    cell = shape.table.cell(r, c); cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor.from_string(INK if r == 0 else SURFACE if r % 2 else 'EDF1F4')
                    cell.margin_left = cell.margin_right = Inches(.17); cell.margin_top = cell.margin_bottom = Inches(.07)
                    sample = ['観点', '確認事項', '留保'][c] if r == 0 else [SAMPLE_HEADS[r-1],
                        ['索引と用語の対応を確認', '引用箇所と調査条件を照合', '担当と確認時期を決める', '適用できない場面を記録', '評価項目の定義を揃える', '未確認の論点を引き継ぐ'][r-1],
                        ['架空資料A・検索節\n部署間の呼称差が残る', '架空資料B・条件節\n原文の条件を保持', '架空資料C・運用節\n確認日：未記録', '架空資料D・例外節\n個別の判断が必要', '資料間の定義に差がある', '確認先の合意が必要'][r-1]][c]
                    style(cell.text_frame, sample, 18, 'FFFFFF' if r == 0 else INK, r == 0)
                    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                    b.entry['texts'][f'cell_{r}_{c}'] = dict(shape=shape.name, cell=[r, c], sample=sample, role='table', vertical_anchor='middle', **capacity(12 / cols - .30, 3.9 / (n + 1) - .1, 18))
                b.entry['rows'][f'row_{r}']={'top_emu':Inches(2.35+r*3.9/(n+1)), 'height_emu':Inches(3.9/(n+1)),
                    'default_alignment':'middle','members':[{'shape':shape.name,'cell':[r,c]} for c in range(cols)]}
            b.slot('note', '留保：同じ条件で比較した結果ではない。各資料の対象範囲を確認する。', .72, 6.35, 11.8, .43, 14, MUTED, role='evidence')
        elif f == 'chart':
            side = v.startswith('bar'); cx = .66 if v == 'bar_left' else 5.0 if side else .66
            cy = 2.73 if side else 2.71 if v != 'column_bottom' else 4.03
            cw, ch = (7.64, 3.77) if side else (12, 2.42)
            b.slot('unit', '課題の言及数（件）／複数回答・架空例', cx, cy-.37, cw, .32, 14, MUTED, role='evidence')
            data = CategoryChartData(); data.categories = ['検索', '出典', '更新', '共有']; data.add_series('架空集計（件）', [42, 35, 28, 18])
            shape = b.slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED if side else XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(cx), Inches(cy), Inches(cw), Inches(ch), data); shape.name = 'slot:chart'
            chart = shape.chart; chart.has_legend = False; chart.has_title = False
            chart.value_axis.minimum_scale = 0
            for axis in (chart.category_axis, chart.value_axis):
                axis.tick_labels.font.name = FONT; axis.tick_labels.font.size = Pt(15)
                axis.tick_labels.font.color.rgb = RGBColor.from_string(MUTED)
                axis.format.line.fill.background()
                axis.major_tick_mark = XL_TICK_MARK.NONE; axis.minor_tick_mark = XL_TICK_MARK.NONE
                axis.tick_labels.number_format = 'General'
            plot = chart.plots[0]; plot.has_data_labels = True; plot.data_labels.position = XL_DATA_LABEL_POSITION.OUTSIDE_END
            plot.data_labels.font.size = Pt(16); plot.data_labels.font.name = FONT; plot.data_labels.number_format = 'General'
            plot.data_labels.font.color.rgb = RGBColor.from_string(INK)
            chart.value_axis.major_gridlines.format.line.color.rgb = RGBColor.from_string(LINE)
            chart.value_axis.major_gridlines.format.line.width = Pt(.6)
            chart.series[0].format.fill.solid(); chart.series[0].format.fill.fore_color.rgb = RGBColor.from_string(TEAL); chart.series[0].format.line.fill.background()
            b.entry['charts']['chart'] = dict(shape=shape.name, kind='category', category_count=4, series_count=1, point_counts=[4], max_label_units=10, domain=[0, 1e12],
                sample={'categories': list(data.categories), 'series': [{'name': '架空集計（件）', 'values': [42, 35, 28, 18], 'x_values': []}]})
            b.entry['charts']['chart']['sample']['categories'] = ['検索', '出典', '更新', '共有']
            tx, ty, tw, th = (8.7 if v == 'bar_left' else .66, 2.52, 3.96, 3.1) if side else (.8, 5.45 if v == 'column_top' else 2.33, 11.7, .75)
            b.slot('insight', '検索の課題が多いが、回答件数のみで重要度を判断することはできない。', tx, ty, tw, th, 21, bold=True)
            b.slot('note', '架空集計であり、効果や因果関係は未検証。', tx, 5.86 if side else 6.36 if v == 'column_top' else 3.17, tw, .72 if side else .40, 14, MUTED, role='evidence')
        elif f == 'quote':
            if v == 'top': qx, qw, tx, tw, ty = .66, 12, .76, 11.8, 5.08
            else: qx, qw, tx, tw, ty = (.66, 6.2, 7.35, 5.3, 2.55) if v == 'left' else (6.46, 6.2, .66, 5.3, 2.55)
            qh = 2.4 if v == 'top' else 4.22
            b.rect(qx, 2.34, qw, qh, TINT)
            b.slot('quotation', '資料を見つけた後で、いま使える内容かを判断するのに時間がかかる。', qx + .24, 2.71, qw - .48, qh - 1.03, 25, TEAL, True, 'quotation')
            b.slot('attribution', '架空ヒアリングの発言例。一般化は未検証。', qx + .3, 2.34 + qh - .66, qw - .6, .52, 14, MUTED, role='evidence')
            b.slot('interpretation', '示唆：検索結果と一緒に更新日・適用範囲・確認担当を示す。', tx, ty, tw, 1.05 if v == 'top' else 2.6, 21)
            b.slot('caveat', '留保：発言と分析者の解釈を区別する。', tx, 6.23, tw, .45, 14, MUTED, role='evidence')
            b.entry['text_flow'].extend([{'body':'quotation','following':'attribution','gap_emu':Inches(.16)},
                                         {'body':'interpretation','following':'caveat','gap_emu':Inches(.16)}])
    e = b.entry
    # Slots have stable semantic identities across same-family variants.
    e['slot_mapping'] = {k: v['shape'] + (':' + ','.join(map(str, v['cell'])) if 'cell' in v else '') for k, v in e['texts'].items()}
    e['emphasis_item'] = 'item_1' if (f=='cards' and v.startswith('feature')) or (f=='summary' and v=='split') else None
    e['semantic_signature'] = {kind: sorted(e[kind]) for kind in ('texts', 'charts', 'metrics', 'states', 'images')}
    if f == 'cards':
        visual = 'feature_with_support' if v.startswith('feature') else 'image_with_cards' if v.startswith('image') else 'equal_columns' if v=='grid' and n<=3 else 'card_grid' if v=='grid' else 'editorial_rows'
    elif f in ('bullets', 'agenda', 'summary') and v=='columns': visual = 'equal_columns' if n<=3 else 'card_grid'
    elif f == 'metrics' and v=='row' or f in ('steps','timeline') and v=='horizontal': visual = 'equal_columns'
    elif f in ('diagram','quote'): visual = f + ('_horizontal' if v != 'top' else '_top')
    elif f=='comparison' and v in ('columns','swapped'): visual = 'comparison_columns'
    elif f=='takeaway' and v in ('left','right'): visual = 'takeaway_split'
    elif f=='chart': visual = 'chart_side' if v.startswith('bar') else 'chart_stacked'
    else: visual = f'{f}:{v}'
    e['visual_signature'] = visual
    e['selection_intent'] = f'{e["name"]} / {n} items / {v}; preserve semantic slot identity and source order'
    e['split_policy'] = 'reject_overflow; explicit replan with retained origin and references'
    e['image_policy'] = 'named asset slot; explicit fit/crop, no network' if e['images'] else 'no image slots; select image cards or legacy text_image for supplied photos'
    e['capacity'] = dict(text_slots=len(e['texts']), chart_slots=len(e['charts']), metric_slots=len(e['metrics']), state_slots=0,
                         item_count=n, required='exact semantic keys; fixed geometry and sizes; no shrink')
    e['template'] = f'catalog/editorial/templates/{e["layout_id"]}.pptx'
    e['preview'] = f'catalog/editorial/previews/{e["layout_id"]}.png'
    e['schema'] = f'catalog/editorial/schemas/{e["layout_id"]}.schema.json'
    e['shape_count'] = len(b.slide.shapes); e['native_table_count'] = sum(s.has_table for s in b.slide.shapes); e['native_chart_count'] = sum(s.has_chart for s in b.slide.shapes)
    b.slide.notes_slide.notes_text_frame.text = json.dumps({'layout_id': e['layout_id'], 'license': (ROOT / 'LICENSE').read_text(encoding='utf-8'), 'provenance': 'Original design; fictional content'}, ensure_ascii=False)
    b.prs.core_properties.title = e['name']; b.prs.core_properties.author = 'slide-layout-agent'
    destination = ROOT / e['template']; destination.parent.mkdir(parents=True, exist_ok=True); b.prs.save(destination)
    e['template_sha256'] = hashlib.sha256(destination.read_bytes()).hexdigest()
    return e


def main():
    entries = [build_one(spec) for spec in SPECS]
    write(OUT / 'manifest.json', {'version': 1, 'catalog': 'editorial', 'layout_count': len(entries), 'family_count': 14, 'layouts': entries})
    from slide_agent.catalog import editorial_registry, all_registry
    editorial_registry.cache_clear()
    from slide_agent.schema_contract import layout_schema
    from slide_agent.models import Plan
    for e in all_registry().values(): write(ROOT / e.get('schema', f'catalog/schemas/{e["layout_id"]}.schema.json'), layout_schema(e))
    write(ROOT / 'catalog/plan.schema.json', Plan.model_json_schema())
    print(f'Built {len(entries)} original templates across 14 families')


if __name__ == '__main__': main()
