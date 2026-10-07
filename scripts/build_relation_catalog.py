"""Build bounded, original native diagrams. Runtime consumes the saved PPTX files."""
import hashlib
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt
from scripts.build_editorial_catalog import Builder, capacity, style, write, INK, MUTED, TEAL, LINE, SURFACE, TINT, BG
from slide_agent.relation_specs import SPECS

WIDTH = 13.333333


def box(b, name, bounds, text='主体', size=18, bold=False):
    x,y,w,h = bounds
    s=b.text(name,text,x,y,w,h,size,INK,bold)
    s.text_frame.vertical_anchor=MSO_ANCHOR.MIDDLE
    return {'shape':name,'bounds_emu':[s.left,s.top,s.width,s.height],**capacity(w,h,size)}


def line(b, name, points, arrow=True, muted=False):
    x1,y1,x2,y2=points
    s=b.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,*[Inches(v) for v in points]);s.name=name
    s.line.color.rgb=RGBColor.from_string(LINE if muted else TEAL);s.line.width=Pt(.8 if muted else 1.7)
    s._element.spPr.append(OxmlElement('a:effectLst'))
    if arrow:
        end=OxmlElement('a:tailEnd');end.set('type','triangle');end.set('w','med');end.set('len','med')
        s._element.spPr.get_or_add_ln().append(end)
    return s


def build(spec):
    b=Builder({**spec,'name':f"Native {spec['kind']} / {spec['item_count']}", 'structural_variant':f"{spec['direction']}_{spec['max_events']}"})
    e=b.entry;e['catalog']='relations'
    b.slot('title','主体と通信の関係を確認する',.66,.69,12,1.08,28,bold=True,role='title')
    b.slot('context','架空の受付・照会システム。矢印と囲みは異なる関係を表す。',.68,1.85,11.95,.34,15,MUTED,role='context')
    b.slot('caveat','注記：架空の構造例。実在の仕様や安全性を示すものではない。',.68,6.49,11.95,.33,14,MUTED,role='evidence')
    kind,n=spec['kind'],spec['item_count']
    mirror=spec['direction']=='rl'
    reflect=lambda bounds:[WIDTH-bounds[0]-bounds[2],*bounds[1:]] if mirror else bounds
    flipline=lambda p:[WIDTH-p[0],p[1],WIDTH-p[2],p[3]] if mirror else p
    if kind=='compare':
        shape=b.slide.shapes.add_table(4,4,Inches(.66),Inches(2.4),Inches(12),Inches(3.84));shape.name='rel:comparison'
        table=shape.table
        for col,w in zip(table.columns,[2.34,3.22,3.22,3.22]):col.width=Inches(w)
        table.rows[0].height=Inches(.66)
        for row in list(table.rows)[1:]:row.height=Inches(1.06)
        cells={}
        for r in range(4):
            for c in range(4):
                cell=table.cell(r,c);cell.margin_left=cell.margin_right=Inches(.12);cell.margin_top=cell.margin_bottom=Inches(.05)
                cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string(INK if r==0 else SURFACE if r%2 else 'EDF2F6')
                size=18 if c or not r else 17
                style(cell.text_frame,'比較軸' if r==c==0 else '方式' if not r else '説明',size,'FFFFFF' if not r else INK,r==0 or c==0)
                cell.vertical_anchor=MSO_ANCHOR.MIDDLE
                cells[f'{r}:{c}']={'shape':shape.name,'cell':[r,c],**capacity(table.columns[c].width/Inches(1)-.24,table.rows[r].height/Inches(1)-.10,size)}
        e['relation']={'kind':kind,'cells':cells}
    else:
        # Only the registered boundary is dynamic; it is never an arrow endpoint.
        boundary=b.rect(.7,2.25,3.3,4.05,TINT,LINE);boundary.name='rel:boundary'
        boundary_label=box(b,'rel:boundary_label',[.8,2.28,3.1,.31],'内包：管理主体',14,True)
        if kind=='fan':
            positions=[[.9,2.76,3.1,3.58]]+[[9.27,2.76+i*3.58/(n-1),3.1,3.58/(n-1)-.16] for i in range(n-1)]
            pairs=[(0,i) for i in range(1,n)]
        elif kind=='pair':
            positions=[[.9,3.56,3.1,1.35],[9.27,3.56,3.1,1.35]];pairs=[(0,1)]
        elif kind=='spoke':
            positions=[[5.13,4.58,3.05,1.12],[.8,4.58,3.05,1.12],[9.48,4.58,3.05,1.12],[5.13,2.64,3.05,.85]];pairs=[(0,i) for i in range(1,4)]
        elif kind=='pair_fan':
            b.rect(6.55,2.35,.012,3.98,LINE)
            b.slot('left_caption','一対一の照会',*reflect([.8,2.28,5.5,.4]),20,TEAL,True)
            b.slot('right_caption','一対多の配信',*reflect([6.85,2.28,5.5,.4]),20,TEAL,True)
            positions=[[.85,3.64,1.7,1.18],[4.57,3.64,1.7,1.18],[6.97,3.1,1.7,2.82],[10.8,3.1,1.7,.82],[10.8,5.1,1.7,.82]];pairs=[(0,1),(2,3),(2,4)]
        else:
            width=2.13 if n==5 else 2.45
            centers=[1.53+i*10.26/(n-1) for i in range(n)]
            positions=[[x-width/2,2.61,width,.54] for x in centers];pairs=[]
        nodes=[]
        for i,pos in enumerate(positions):
            p=reflect(pos);shape=b.rect(*p,SURFACE,TEAL);shape.name=f'rel:node_{i}'
            x,y,w,h=p
            label=box(b,f'rel:node_label_{i}',[x+.08,y+.025,w-.16,h-.05],size=18,bold=True)
            # Titles inside a tall hub are still short labels, never paragraphs.
            label['max_lines']=min(2,label['max_lines'])
            nodes.append({'shape':shape.name,'bounds_emu':[shape.left,shape.top,shape.width,shape.height],'label':label})
        groups=[[i] for i in range(n)]
        if kind=='fan':groups += [list(range(1,n))]
        if kind=='sequence':groups=[list(range(a,z)) for a in range(n) for z in range(a+1,n+1)]
        if kind=='pair_fan':groups=[]
        relation={'kind':kind,'nodes':nodes,'boundary_label':boundary_label,'allowed_boundaries':groups,'routes':{}}
        if kind=='sequence':
            maximum=spec['max_events'];step=3.18/maximum
            relation.update(max_events=maximum,event_top=Inches(3.25),event_step=Inches(step),label_height=Inches(.59 if maximum==4 else .31),label_font=16,
                            node_centers=[p['bounds_emu'][0]+p['bounds_emu'][2]//2 for p in nodes])
            for i,node in enumerate(nodes):
                x=relation['node_centers'][i]/Inches(1)
                line(b,f'rel:lifeline_{i}',[x,3.15,x,6.41],False,True)
            for i in range(maximum):
                y=3.25+i*step
                label=box(b,f'rel:event_label_{i}',[.8,y,11.7,.59 if maximum==4 else .31],size=16)
                label_shape=b.slide.shapes[-1]
                label_shape.fill.solid();label_shape.fill.fore_color.rgb=RGBColor.from_string(BG)
                label['max_lines']=2 if maximum==4 else 1
                relation.setdefault('event_labels',[]).append(label)
                for part in range(3):line(b,f'rel:event_{i}_{part}',[1, y+.34,12,y+.34],part==2)
        else:
            for a,z in pairs:
                for source,target in ((a,z),(z,a)):
                    start,end=positions[source],positions[target]
                    if kind=='spoke' and {source,target}=={0,3}:
                        cx=positions[0][0]+positions[0][2]/2+(.18 if source==3 else -.18)
                        points=[cx,start[1]+(start[3] if source==3 else 0),cx,end[1]+(0 if target==0 else end[3])]
                        labelbounds=[cx+.16 if source==3 else cx-1.46,3.75,1.3,.55]
                    else:
                        left=start[0]<end[0]
                        cy=(positions[z][1]+positions[z][3]/2 if kind=='fan' or kind=='pair_fan' and a==2 else start[1]+start[3]/2)
                        cy+=-.20 if source==a else .20
                        points=[start[0]+(start[2] if left else 0),cy,end[0]+(0 if left else end[2]),cy]
                        lx=min(points[0],points[2]);w=abs(points[2]-points[0])
                        labelbounds=[lx+.08,cy-.30,w-.16,.29]
                    key=f'{source}>{target}'
                    line(b,'rel:edge_'+key,flipline(points))
                    label=box(b,'rel:edge_label_'+key,reflect(labelbounds),size=15)
                    label['max_lines']=min(label['max_lines'],2 if kind=='spoke' and {source,target}=={0,3} else 1)
                    relation['routes'][key]={'shape':'rel:edge_'+key,'label':label,'points_emu':[Inches(v) for v in flipline(points)]}
        e['relation']=relation
    e.update(semantic_signature={'kind':kind,'nodes':n},visual_signature=f'relations:{kind}',slot_mapping={'network':'stable node IDs and directed edge IDs; ordering retained','comparison':'methods by ID; values aligned to shared axes'},
             selection_intent='Preserve directed endpoints, chronological order, boundaries and shared references; mirror only spatial positions.',
             split_policy='reject_overflow; explicitly split with source references',image_policy='native only; no image slots',
             capacity={'item_count':n,'title_lines':2,'max_events':spec['max_events'],'required':'registered topology and bounds; no arbitrary coordinates; no shrink'})
    e['template']=f"catalog/relations/templates/{e['layout_id']}.pptx";e['schema']=f"catalog/relations/schemas/{e['layout_id']}.schema.json";e['preview']=f"catalog/relations/previews/{e['layout_id']}.png"
    e['shape_count']=len(b.slide.shapes)
    b.prs.core_properties.title=e['name'];b.prs.core_properties.author='slide-layout-agent'
    p=ROOT/e['template'];p.parent.mkdir(parents=True,exist_ok=True);b.prs.save(p)
    e['template_sha256']=hashlib.sha256(p.read_bytes()).hexdigest()
    return e


def main():
    entries=[build(s) for s in SPECS]
    write(ROOT/'catalog/relations/manifest.json',{'version':1,'catalog':'relations','layouts':entries})
    from slide_agent.catalog import relation_registry
    relation_registry.cache_clear()
    from slide_agent.schema_contract import layout_schema
    from slide_agent.models import Plan
    for e in entries:write(ROOT/e['schema'],layout_schema(e))
    write(ROOT/'catalog/plan.schema.json',Plan.model_json_schema())
    print(f'Built {len(entries)} native relationship/communication templates')


if __name__=='__main__':main()
