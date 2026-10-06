"""Public fictional reproduction deck for the independent review regressions."""
import copy
import json
import sys
from pathlib import Path

ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
from scripts.make_catalog_samples import sample_plan
from slide_agent.catalog import registry
from slide_agent.models import Theme
from slide_agent.validate import validate
from slide_agent.render import render
from slide_agent.audit import audit
from pptx import Presentation


def main():
    entries=[];cases=[]
    for variant in ('warm','cool'):
        cagr=copy.deepcopy(registry()['scenario_lines_cagr'])
        cagr['title']['sample']='10年間の年平均成長率（架空例）'
        chart=next(iter(cagr['charts'].values()))['sample']
        chart['categories']=[str(y) for y in range(2020,2031,2)];chart['elapsed_years']=10.
        cagr['texts']['text_012']['sample']='単位、2020〜2030年'
        entries.append(cagr);cases.append(('cagr_ten_years',variant))
        for case,values,title in [('bridge_sign_flip',[11,3,-4,3,13],'増減の符号を反転（架空例）'),
                                  ('decimal_precision',[123456.7,1,2,3,123462.7],'入力小数を保持（架空例）')]:
            e=copy.deepcopy(registry()['true_waterfall']);e['title']['sample']=title
            for k,v in zip(e['bridge']['keys'],values):e['metrics'][k]['sample']=v
            entries.append(e);cases.append((case,variant))
    source,plan=sample_plan(entries)
    for slide,(_,variant) in zip(plan.slides,cases):slide.variant=variant
    out=ROOT/'catalog/qa/review-regressions';out.mkdir(parents=True,exist_ok=True)
    report=validate(plan,source,out,Theme())
    if not report['ok']:raise ValueError(report['issues'])
    for name,obj in [('source',source),('plan',plan)]:
        (out/f'cases.{name}.json').write_text(obj.model_dump_json(indent=2),encoding='utf-8')
    dest=out/'cases.pptx';render(plan,source,out,Theme(),dest);report['ooxml']=audit(dest,Theme())
    if not report['ooxml']['ok']:raise ValueError(report['ooxml'])
    results=[]
    for i,(slide,e,(case,variant)) in enumerate(zip(Presentation(dest).slides,entries,cases),1):
        shapes={s.name:s for s in slide.shapes}
        actual={'case':case,'variant':variant,'slide':i}
        if case=='cagr_ten_years':
            actual.update(period=shapes['shape:017'].text,callouts=[shapes[f'shape:{j:03}'].text for j in (19,21,23)])
        else:
            actual['bars']=[{'input':e['metrics'][key]['sample'],'label':shapes[label].text,
                             'fill':str(shapes[bar].fill.fore_color.rgb),'top':shapes[bar].top,'height':shapes[bar].height}
                            for key,label,bar in zip(e['bridge']['keys'],e['bridge']['labels'],e['bridge']['shapes'])]
        results.append(actual)
    report['reproductions']=results
    (out/'results.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(results,ensure_ascii=False,indent=2))


if __name__=='__main__':main()
