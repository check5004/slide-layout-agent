"""Publish the actual old/new PowerPoint renders; verify identical content first."""
import argparse
import hashlib
import html
import json
import shutil
import sys
from pathlib import Path

ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
from pptx import Presentation
from scripts.publish_editorial_assets import contact


def values(path):
    result=[]
    for slide in Presentation(path).slides:
        fields={}
        for shape in slide.shapes:
            if shape.has_text_frame:fields[shape.name]=shape.text
            if shape.has_table:
                fields[shape.name]=[[c.text for c in row.cells] for row in shape.table.rows]
        result.append(fields)
    return result


def main():
    parser=argparse.ArgumentParser();parser.add_argument('--rendered',type=Path,required=True)
    parser.add_argument('--out',type=Path,default=ROOT/'examples/editorial-rows');args=parser.parse_args()
    dest=args.out;plan=json.loads((dest/'rows.after.plan.json').read_text(encoding='utf-8'))
    before=json.loads((dest/'rows.before.plan.json').read_text(encoding='utf-8'))
    original=json.loads(json.dumps(plan))
    for slide in original['slides']:slide.pop('row_alignments',None)
    if original!=before:raise ValueError('before/after must retain all source content and refs')
    if values(dest/'rows.before.pptx')!=values(dest/'rows.after.pptx'):raise ValueError('visible native text changed')
    reports={}
    for version in ('before','after'):
        deck=dest/f'rows.{version}.pptx';report=json.loads((args.rendered/f'{version}-render/render-report.json').read_text(encoding='utf-8-sig'))
        if report['pptxSha256']!=hashlib.sha256(deck.read_bytes()).hexdigest():raise ValueError('render SHA mismatch')
        if report['slideCount']!=len(plan['slides']) or report['overflowCount'] or report.get('positionOverflowCount',0):raise ValueError('real render overflow or missing page')
        reports[version]=report
        shutil.copyfile(args.rendered/f'{version}-render/render-report.json',dest/f'{version}.render-report.json')
    images=[];labels=[];cards=[];coverage=[]
    notes={v:[json.loads(s.notes_slide.notes_text_frame.text) for s in Presentation(dest/f'rows.{v}.pptx').slides] for v in ('before','after')}
    (dest/'previews').mkdir(exist_ok=True)
    for index,slide in enumerate(plan['slides'],1):
        pair=[];record={'page':index,'layout_id':slide['layout_id'],'row_alignments':slide['row_alignments']}
        for version in ('before','after'):
            target=dest/'previews'/f'{version}-{index:02}.png'
            shutil.copyfile(args.rendered/f'{version}-render/slide-{index:02}.png',target)
            images.append(target);labels.append(f'{index:02} {version}: {slide["layout_id"]}')
            record[version]={'preview':f'previews/{target.name}','png_sha256':hashlib.sha256(target.read_bytes()).hexdigest(),
                             'template_sha256':notes[version][index-1]['template_sha256']}
            pair.append(f'<figure><figcaption>{version}</figcaption><a href="previews/{target.name}"><img src="previews/{target.name}" alt="{version}: {html.escape(slide["title"]["text"])}"></a></figure>')
        coverage.append(record)
        cards.append(f'<article><h2>{index:02} {html.escape(slide["title"]["text"])}</h2><code>{slide["layout_id"]}</code><div class="pair">{"".join(pair)}</div></article>')
    contact(images,labels,dest/'contact-sheet.png',columns=2)
    for start in range(0,len(images),8):contact(images[start:start+8],labels[start:start+8],dest/f'contact-{start//8+1:02}.png',columns=2)
    (dest/'index.html').write_text('''<!doctype html><html lang="ja"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>行の上下配置 before / after</title><style>
body{margin:0;background:#eef2f5;color:#213547;font:16px/1.6 system-ui,sans-serif}header{background:#213547;color:white;padding:28px 4vw}header p{max-width:980px}main{padding:24px 3vw}article{padding:20px;background:white;border:1px solid #dce3e9;margin:0 0 24px;border-radius:8px}h1{margin:0;font-size:28px}h2{margin:0;font-size:20px}.pair{display:grid;grid-template-columns:1fr 1fr;gap:18px}figure{margin:12px 0 0}img{display:block;width:100%}figcaption,code{color:#087e83}a{color:#087e83}@media(max-width:850px){.pair{grid-template-columns:1fr}}</style>
<header><h1>行の上下配置 / before → after</h1><p>同じ文字・数値・引用・項目順。2〜6項目、1行・複数行、本文と根拠を混在させた実描画です。短いラベルは行の中央、本文＋根拠は一つのまとまりとして配置します。11枚目は行別指定、12枚目は明示的な上寄せ。文章の読み順が重要なカード内の縦積みは上寄せを維持します。</p></header><main>'''+''.join(cards)+'</main></html>',encoding='utf-8')
    (dest/'comparison.json').write_text(json.dumps({'before_commit':'23f6cb8ffea5ec8d6b97ae61f13759feb3d75e86',
        'before_pptx_sha256':reports['before']['pptxSha256'],'after_pptx_sha256':reports['after']['pptxSha256'],
        'same_source_content_refs_and_order':True,'same_native_visible_text':True,'pages':coverage},ensure_ascii=False,indent=2),encoding='utf-8')
    print('Published 12 verified before/after pairs; native content and reference equality passed')


if __name__=='__main__':main()
