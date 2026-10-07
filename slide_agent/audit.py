"""Inspect generated geometry and editable OOXML parts, without claiming rendering."""
from pathlib import Path
import json
from zipfile import ZipFile

from lxml import etree
from pptx import Presentation
from pptx.util import Inches
from .catalog import is_catalog, all_registry as registry, template_path


def audit(path, theme):
    prs = Presentation(str(path))
    issues, counts = [], {"slides": len(prs.slides), "text_shapes": 0, "tables": 0, "charts": 0, "pictures": 0}
    allowed_links=set()
    for i, slide in enumerate(prs.slides, 1):
        try:
            note = json.loads(slide.notes_slide.notes_text_frame.text)
        except (ValueError, AttributeError):
            note = {}
        catalog = is_catalog(note.get("layout_id"))
        if note.get('display'):
            from .models import DeckDisplay
            from .reader_display import parts,expected_urls
            try:
                display=DeckDisplay.model_validate(note['display']);source_ids=note.get('display_source_ids',[])
                urls=expected_urls(display,source_ids);allowed_links.update(urls)
                if display.mode=='reader':
                    footer=next((s for s in slide.shapes if s.name==('meta:footer' if catalog else 'footer')),None)
                    if footer is None or footer.text!=''.join(t for t,u in parts(display,source_ids)):
                        issues.append(f'slide {i}: reader source footer changed or missing')
                    if footer and {r.hyperlink.address for p in footer.text_frame.paragraphs for r in p.runs if r.hyperlink.address}!=urls:
                        issues.append(f'slide {i}: reader source links changed or missing')
                    if any(s.text for s in slide.shapes if s.name in ('meta:layout','meta:format')):
                        issues.append(f'slide {i}: internal QA metadata leaked into reader mode')
            except (ValueError,TypeError):issues.append(f'slide {i}: invalid reader display metadata')
        if catalog:
            entry = registry()[note["layout_id"]]
            overrides = note.get('row_alignments', {})
            if not isinstance(overrides,dict) or any(k not in entry.get('rows',{}) or v not in ('top','middle') for k,v in overrides.items()):
                issues.append(f'slide {i}: invalid registered row alignment')
                overrides = {}
            original = Presentation(str(template_path(entry))).slides[0]
            if entry.get('catalog')=='relations':
                from .models import RelationContents
                from .relations import populate
                try:populate(original,entry,RelationContents.model_validate(note.get('relation_contents')))
                except (ValueError,TypeError,KeyError):issues.append(f'slide {i}: missing or invalid relationship plan in notes')
            expected = {s.name: s for s in original.shapes}
            actual = {s.name: s for s in slide.shapes}
            if entry.get('catalog')=='relations':
                from .relations import audit_semantics
                issues.extend(f'slide {i}: {message}' for message in audit_semantics(expected,actual))
            if set(expected) != set(actual) or len(actual) != len(slide.shapes):
                issues.append(f"slide {i}: native template shape identity changed")
            dynamic = {m.get("shape") for m in entry["metrics"].values() if m.get("shape")}
            for flow in entry.get('text_flow', []):
                bn=entry['texts'][flow['body']]['shape'];fn=entry['texts'][flow['following']]['shape']
                if bn not in actual or fn not in actual:
                    continue  # The identity check above already reports the missing slot.
                dynamic.update([bn,fn]);body,following=actual[bn],actual[fn];ob,of=expected[bn],expected[fn]
                if flow.get('row'):
                    row=entry['rows'][flow['row']];end=row['top_emu']+row['height_emu']
                    total=body.height+flow['gap_emu']+following.height
                    expected_top=row['top_emu']+((row['height_emu']-total)//2 if overrides.get(flow['row'],row['default_alignment'])=='middle' else 0)
                    if (body.left,body.width)!=(ob.left,ob.width) or not 0<body.height<=ob.height or body.top!=expected_top or body.top<row['top_emu']:
                        issues.append(f'slide {i}: text flow body escaped its registered row alignment')
                    if (following.left,following.width,following.height)!=(of.left,of.width,of.height) or following.top!=body.top+body.height+flow['gap_emu'] or following.top+following.height>end:
                        issues.append(f'slide {i}: text flow caveat escaped its registered row group')
                else:
                    if (body.left,body.top,body.width)!=(ob.left,ob.top,ob.width) or not 0<body.height<=ob.height:
                        issues.append(f'slide {i}: text flow body escaped its registered bounds')
                    if (following.left,following.width,following.height)!=(of.left,of.width,of.height) or not body.top+body.height <= following.top <= of.top:
                        issues.append(f'slide {i}: text flow caveat escaped its registered bounds')
            for spec in entry.get('images', {}).values():
                dynamic.add(spec['shape'])
                pic = actual.get(spec['shape'])
                x, y, w, h = spec['bounds_emu']
                if pic is None or pic.shape_type != 13 or pic.left < x-2 or pic.top < y-2 or pic.left+pic.width > x+w+2 or pic.top+pic.height > y+h+2:
                    issues.append(f'slide {i}: image escaped or missing from named slot {spec["shape"]}')
            if entry.get("bridge"):
                dynamic.update(entry["bridge"]["shapes"]+entry["bridge"]["labels"])
            if entry.get("calculation"):
                dynamic.update(entry["calculation"]["totals"])
                for group in entry["calculation"]["groups"]:
                    for bar in group:
                        dynamic.update([bar["shape"],bar["label"],bar.get("extra_shape")])
            if entry.get("reference_line"):
                dynamic.update([entry["reference_line"]["shape"],entry["reference_line"]["label"]])
            for name in set(expected)&set(actual)-dynamic:
                a,b=expected[name],actual[name]
                if (a.left,a.top,a.width,a.height)!=(b.left,b.top,b.width,b.height):
                    issues.append(f"slide {i}: fixed template geometry changed: {name}")
            for name in set(expected)&set(actual):
                a,b=expected[name],actual[name]
                oldframes=[c.text_frame for row in a.table.rows for c in row.cells] if a.has_table else [a.text_frame] if a.has_text_frame else []
                newframes=[c.text_frame for row in b.table.rows for c in row.cells] if b.has_table else [b.text_frame] if b.has_text_frame else []
                for oldtf,newtf in zip(oldframes,newframes):
                    oldsizes=[r.font.size.pt for p in oldtf.paragraphs for r in p.runs if r.font.size]
                    newsizes=[r.font.size.pt for p in newtf.paragraphs for r in p.runs if r.font.size]
                    if oldsizes and any(size<min(oldsizes)-.01 for size in newsizes):
                        issues.append(f"slide {i}: font shrunk below template: {name}")
            if entry.get('catalog') == 'editorial':
                from .vertical import ANCHORS,anchor_target,row_anchors
                expected_anchors={}
                specs=[entry['title'],*entry['texts'].values(),*[m['label_capacity'] for m in entry['metrics'].values() if m.get('label_capacity')]]
                for spec in specs:
                    key=(spec['shape'],tuple(spec.get('cell',[])))
                    expected_anchors[key]=(spec,spec.get('vertical_anchor','top'))
                for member,anchor in row_anchors(entry,overrides):
                    expected_anchors[(member['shape'],tuple(member.get('cell',[])))]=(member,anchor)
                for member,anchor in expected_anchors.values():
                    if member['shape'] in actual and anchor_target(actual,member).vertical_anchor!=ANCHORS[anchor]:
                        issues.append(f'slide {i}: native vertical anchor changed: {member["shape"]}')
                occupied = [s for s in slide.shapes if s.name.startswith('slot:')]
                for j, a in enumerate(occupied):
                    for b in occupied[j+1:]:
                        ox = min(a.left+a.width,b.left+b.width)-max(a.left,b.left)
                        oy = min(a.top+a.height,b.top+b.height)-max(a.top,b.top)
                        if ox > 2 and oy > 2:
                            issues.append(f'slide {i}: editorial semantic slots overlap: {a.name}/{b.name}')
        content_shapes = 0
        text_boxes = []
        for shape in slide.shapes:
            if min(shape.left, shape.top) < -2 or shape.left + shape.width > prs.slide_width + 2 or shape.top + shape.height > prs.slide_height + 2:
                issues.append(f"slide {i}: out-of-bounds shape {shape.name}")
            if shape.has_chart:
                counts["charts"] += 1
                content_shapes += 1
            if shape.has_table:
                counts["tables"] += 1
                content_shapes += 1
                frames = [(cell.text_frame, 6 if catalog else theme.table_pt) for row in shape.table.rows for cell in row.cells]
            elif shape.has_text_frame and shape.text:
                if shape.name in ("body", "title", "lead", "label", "footer"):
                    text_boxes.append(shape)
                counts["text_shapes"] += 1
                content_shapes += int(catalog or shape.name in ("body", "title"))
                minimum = 6 if catalog else theme.title_pt if shape.name == "title" else theme.body_pt if shape.name == "body" else 11
                frames = [(shape.text_frame, minimum)]
            else:
                frames = []
            if shape.shape_type == 13:
                counts["pictures"] += 1
            for tf, minimum in frames:
                for p in tf.paragraphs:
                    for run in p.runs:
                        if run.font.size is None or run.font.size.pt < minimum:
                            issues.append(f"slide {i}: missing/undersized font in {shape.name}")
        if not content_shapes:
            issues.append(f"slide {i}: no native editable content")
        for a_index, a in enumerate(text_boxes):
            for b in text_boxes[a_index+1:]:
                overlap_x = min(a.left+a.width, b.left+b.width) - max(a.left, b.left)
                overlap_y = min(a.top+a.height, b.top+b.height) - max(a.top, b.top)
                if overlap_x > 2 and overlap_y > 2:
                    issues.append(f"slide {i}: overlapping text boxes {a.name}/{b.name}")
    with ZipFile(path) as archive:
        if archive.testzip():
            issues.append("invalid ZIP CRC")
        for name in archive.namelist():
            if name.endswith((".xml", ".rels")):
                root = etree.fromstring(archive.read(name))
                if root.xpath("//*[local-name()='p' and count(*[local-name()='pPr']) > 1]"):
                    issues.append(f"duplicate paragraph properties in {name}")
                for font_tag in ('ea','cs','latin'):
                    if root.xpath(f"//*[local-name()='rPr' or local-name()='defRPr'][count(*[local-name()='{font_tag}']) > 1]"):
                        issues.append(f'duplicate font {font_tag} declaration in {name}; PowerPoint may reject the package')
                if name.startswith("ppt/slides/slide") and name.endswith(".xml") and root.xpath("//*[local-name()='normAutofit' or local-name()='spAutoFit']"):
                    issues.append(f"auto-resize present in {name}")
                if name.endswith(".rels"):
                    for relation in root.xpath("//*[@TargetMode='External']"):
                        if not (relation.get('Type','').endswith('/hyperlink') and relation.get('Target') in allowed_links):
                            issues.append(f"unexpected external relationship in {name}")
        counts["embedded_workbooks"] = sum(name.startswith("ppt/embeddings/") and name.endswith(".xlsx") for name in archive.namelist())
        if counts["embedded_workbooks"] != counts["charts"]:
            issues.append("native charts must each have an editable embedded workbook")
    return {"ok": not issues, "issues": issues, "counts": counts, "rendering": "not_performed"}
