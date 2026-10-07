"""Explicit reader metadata and inert HTTP(S) source links; never fetch a URL."""
import copy
import datetime
from pptx import Presentation
from pptx.util import Inches
from .catalog import template_path


def parts(display, source_ids):
    result=[(display.footer,None)] if display.footer else []
    for citation in display.citations:
        if not set(citation.source_ids)&set(source_ids):continue
        result.append((citation.title,citation.url))
        detail=' / '.join(v for v in (citation.version,('確認 '+citation.accessed_on) if citation.accessed_on else '') if v)
        if detail:result.append((' ['+detail+']',None))
    # A separator belongs between citation groups, not between a title and metadata.
    output=[]
    for text,url in result:
        if output and not text.startswith(' ['):output.append(('  |  ',None))
        output.append((text,url))
    return output


def validate_display(plan, source, add):
    from .validate import walk_evidence,units
    from .catalog import all_registry,is_catalog
    known={s.id for s in source.segments}
    for citation in plan.display.citations:
        if not set(citation.source_ids)<=known:add('DISPLAY_SOURCE','display citation references an unknown source segment')
        if citation.accessed_on:
            try:datetime.date.fromisoformat(citation.accessed_on)
            except ValueError:add('DISPLAY_DATE','accessed_on must be a real calendar date')
    if plan.display.mode=='qa':return
    for slide in plan.slides:
        refs={r.source_id for t in walk_evidence(slide) for r in t.refs}
        text=''.join(t for t,u in parts(plan.display,refs))
        if not text:continue
        width,size=11.1,11
        if is_catalog(slide.layout_id):
            template=Presentation(template_path(all_registry()[slide.layout_id]))
            slots=[s for s in template.slides[0].shapes if s.name=='meta:footer']
            if not slots:add('DISPLAY_SLOT','this template has no reader footer slot',slide.id);continue
            slot=slots[0];width=slot.width/Inches(1)
            size=max((r.font.size.pt for p in slot.text_frame.paragraphs for r in p.runs if r.font.size),default=11)
        if '\n' in text or units(text)>(width*72-6)/(size*1.08):
            add('DISPLAY_OVERFLOW','source footer exceeds its fixed line; use concise document titles/versions or fewer sources, never shrink',slide.id)


def apply(slide, display, source_ids, catalog=True):
    if display.mode=='qa':return
    from .template_engine import replace_text
    for s in slide.shapes:
        if s.name in ('meta:layout','meta:format'):replace_text(s.text_frame,'')
    footer=next((s for s in slide.shapes if s.name==('meta:footer' if catalog else 'footer')),None)
    if footer is None:return
    tf=footer.text_frame;replace_text(tf,'')
    p=tf.paragraphs[0];first=p.runs[0];rpr=copy.deepcopy(first._r.get_or_add_rPr())
    for i,(text,url) in enumerate(parts(display,source_ids)):
        run=first if i==0 else p.add_run();run.text=text
        if i:
            old=run._r.get_or_add_rPr();run._r.replace(old,copy.deepcopy(rpr))
        if url:run.hyperlink.address=url;run.font.underline=True


def expected_urls(display, source_ids):
    return {url for text,url in parts(display,source_ids) if url} if display.mode=='reader' else set()
