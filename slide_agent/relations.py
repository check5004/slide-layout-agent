"""Populate only the finite routes/lanes registered in the native catalog."""
import math
from pptx.util import Inches
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.dml.color import RGBColor


def bounds_capacity(bounds, size, maximum):
    return {'max_units_per_line':round(max(1,(bounds[2]/Inches(1)*72-6)/(size*1.08)),2),'max_lines':maximum}


def sequence_geometry(entry, network, index):
    r=entry['relation'];edge=network.edges[index];ids=[n.id for n in network.nodes]
    a,z=ids.index(edge.source),ids.index(edge.target);x1,x2=r['node_centers'][a],r['node_centers'][z]
    y=r['event_top']+index*r['event_step'];height=r['label_height']
    line_y=y+height+Inches(.015)
    if a!=z:
        bounds=[min(x1,x2)+Inches(.10),y,abs(x2-x1)-Inches(.20),height]
        parts=[[x1,line_y,x2,line_y]]
    else:
        sign=-1 if a==len(ids)-1 else 1
        if entry['direction']=='rl':sign*=-1
        x3=x1+sign*Inches(.62);bottom=line_y+Inches(.09)
        parts=[[x1,line_y,x3,line_y],[x3,line_y,x3,bottom],[x3,bottom,x1,bottom]]
        bounds=[x1-Inches(.9),y,Inches(1.8),height]
    return bounds,parts


def boundary_geometry(entry, network):
    r=entry['relation'];group=network.boundaries[0];ids=[n.id for n in network.nodes]
    members=[r['nodes'][ids.index(k)]['bounds_emu'] for k in group.members]
    left=min(p[0] for p in members)-Inches(.1);right=max(p[0]+p[2] for p in members)+Inches(.1)
    top=min(p[1] for p in members)-Inches(.37)
    bottom=Inches(6.42) if r['kind']=='sequence' else max(p[1]+p[3] for p in members)+Inches(.08)
    bounds=[left,top,right-left,bottom-top]
    label=[left+Inches(.07),top+Inches(.025),right-left-Inches(.14),Inches(.30)]
    if r['kind']=='spoke' and group.members==[ids[0]]:
        # The upper actor connects vertically. Reserve the group's caption below
        # the client so neither directed line can pass through the caption.
        node=members[0];top=node[1]-Inches(.1);bottom=node[1]+node[3]+Inches(.42)
        bounds=[left,top,right-left,bottom-top]
        label=[left+Inches(.07),node[1]+node[3]+Inches(.06),right-left-Inches(.14),Inches(.30)]
    return bounds,label


def errors(entry, contents):
    from .validate import units
    found=[];r=entry['relation'];kind=r['kind']
    def check(text,spec,path):
        if not text.strip():found.append(('RELATION_EMPTY',path))
        lines=sum(max(1,math.ceil(units(s)/spec['max_units_per_line'])) for s in text.split('\n'))
        if lines>spec['max_lines']:found.append(('RELATION_OVERFLOW',f'{path}: exceeds {spec["max_lines"]} lines x {spec["max_units_per_line"]} full-width units; split, never shrink'))
    if kind=='compare':
        c=contents.comparison
        if c is None or contents.network is not None:return [('RELATION_KIND','three-method comparison required; network not accepted')]
        if len({m.id for m in c.methods})!=3:found.append(('RELATION_ID','method IDs must be unique'))
        for i,axis in enumerate(c.axes):check(axis.text,r['cells'][f'{i+1}:0'],f'axis {i}')
        for i,method in enumerate(c.methods):
            check(method.label.text,r['cells'][f'0:{i+1}'],method.id)
            for j,t in enumerate(method.values):check(t.text,r['cells'][f'{j+1}:{i+1}'],f'{method.id}/{j}')
        return found
    net=contents.network
    if net is None or contents.comparison is not None:return [('RELATION_KIND','network required; comparison not accepted')]
    ids=[n.id for n in net.nodes];edge_ids=[e.id for e in net.edges]
    if len(ids)!=entry['item_count']:return [('RELATION_NODE_COUNT',f'exactly {entry["item_count"]} named nodes required')]
    if len(set(ids))!=len(ids) or len(set(edge_ids))!=len(edge_ids):return [('RELATION_ID','node and edge IDs must each be unique')]
    if any(e.source not in ids or e.target not in ids for e in net.edges):return [('RELATION_ENDPOINT','every endpoint must reference a declared node, never a boundary')]
    for node,spec in zip(net.nodes,r['nodes']):check(node.label.text,spec['label'],node.id)
    if kind=='sequence':
        if len(net.edges)>r['max_events']:return [('RELATION_EVENT_COUNT',f'at most {r["max_events"]} ordered events; split explicitly')]
        for i,edge in enumerate(net.edges):
            bounds,_=sequence_geometry(entry,net,i)
            check(edge.label.text,bounds_capacity(bounds,r['label_font'],2 if r['max_events']==4 else 1),edge.id)
    else:
        pairs=[]
        for edge in net.edges:
            key=f'{ids.index(edge.source)}>{ids.index(edge.target)}';pairs.append(key)
            if key not in r['routes']:found.append(('RELATION_TOPOLOGY',f'{edge.id}: directed pair {key} is not a supported route'))
            else:check(edge.label.text,r['routes'][key]['label'],edge.id)
        if len(pairs)!=len(set(pairs)):found.append(('RELATION_DUPLICATE_EDGE','one edge per directed pair; use sequence for repeated messages'))
        if {v for edge in net.edges for v in (edge.source,edge.target)}!=set(ids):found.append(('RELATION_DISCONNECTED','each node needs a registered relationship; use a matching count'))
    for group in net.boundaries:
        if len(group.members)!=len(set(group.members)) or any(k not in ids for k in group.members):
            found.append(('RELATION_BOUNDARY','boundary members must be unique declared node IDs'));continue
        indices=sorted(ids.index(k) for k in group.members)
        if indices not in r['allowed_boundaries']:found.append(('RELATION_BOUNDARY','unsupported containment span; use one registered, non-overlapping boundary'))
        _,label=boundary_geometry(entry,net)
        prefix='内包：' if group.kind=='contains' else '同管理：'
        check(prefix+group.label.text,bounds_capacity(label,14,1),'boundary label including kind')
    return found


def remove(shape):
    shape._element.getparent().remove(shape._element)


def set_line(shape, points, arrow):
    shape.begin_x,shape.begin_y,shape.end_x,shape.end_y=points
    ln=shape._element.spPr.get_or_add_ln()
    for tag in ('a:headEnd','a:tailEnd'):
        for old in list(ln.findall(qn(tag))):ln.remove(old)
    if arrow:
        end=OxmlElement('a:tailEnd');end.set('type','triangle');end.set('w','med');end.set('len','med');ln.append(end)


def populate(slide, entry, contents):
    from .template_engine import replace_text
    shapes={s.name:s for s in slide.shapes};r=entry['relation']
    problems=errors(entry,contents)
    if problems:raise ValueError(problems)
    if r['kind']=='compare':
        table=shapes['rel:comparison'].table;c=contents.comparison
        for i,t in enumerate(c.axes):replace_text(table.cell(i+1,0).text_frame,t.text)
        for i,m in enumerate(c.methods):
            replace_text(table.cell(0,i+1).text_frame,m.label.text)
            for j,t in enumerate(m.values):replace_text(table.cell(j+1,i+1).text_frame,t.text)
        return
    net=contents.network;ids=[n.id for n in net.nodes]
    for i,node in enumerate(net.nodes):replace_text(shapes[f'rel:node_label_{i}'].text_frame,node.label.text)
    if net.boundaries:
        box,label=boundary_geometry(entry,net);group=net.boundaries[0]
        for name,bounds in [('rel:boundary',box),('rel:boundary_label',label)]:
            s=shapes[name];s.left,s.top,s.width,s.height=bounds
        prefix='内包：' if group.kind=='contains' else '同管理：'
        replace_text(shapes['rel:boundary_label'].text_frame,prefix+group.label.text)
        if group.kind=='same_owner':
            s=shapes['rel:boundary'];s.fill.background()
            dash=OxmlElement('a:prstDash');dash.set('val','dash');s._element.spPr.get_or_add_ln().append(dash)
    else:
        remove(shapes['rel:boundary']);remove(shapes['rel:boundary_label'])
    if r['kind']=='sequence':
        for i in range(r['max_events']):
            label=shapes[f'rel:event_label_{i}']
            if i>=len(net.edges):
                remove(label)
                for p in range(3):remove(shapes[f'rel:event_{i}_{p}'])
                continue
            bounds,parts=sequence_geometry(entry,net,i)
            label.left,label.top,label.width,label.height=bounds
            replace_text(label.text_frame,net.edges[i].label.text)
            for p in range(3):
                s=shapes[f'rel:event_{i}_{p}']
                if p>=len(parts):remove(s)
                else:set_line(s,parts[p],p==len(parts)-1)
    else:
        used={f'{ids.index(e.source)}>{ids.index(e.target)}':e for e in net.edges}
        for key,spec in r['routes'].items():
            if key in used:replace_text(shapes[spec['label']['shape']].text_frame,used[key].label.text)
            else:
                remove(shapes[spec['shape']]);remove(shapes[spec['label']['shape']])


def audit_semantics(expected, actual):
    """Compare direction, labels, grouping and native shape types to replayed plan."""
    issues=[]
    for name in set(expected)&set(actual):
        if not name.startswith('rel:'):continue
        a,b=expected[name],actual[name]
        if a.shape_type!=b.shape_type:issues.append(f'relation shape type changed: {name}');continue
        if a.has_text_frame:
            if a.text!=b.text or a.text_frame.vertical_anchor!=b.text_frame.vertical_anchor:issues.append(f'relation label/anchor changed: {name}')
        if a.has_table:
            if [c.width for c in a.table.columns]!=[c.width for c in b.table.columns] or [r.height for r in a.table.rows]!=[r.height for r in b.table.rows]:issues.append('comparison column/row geometry changed')
            for ra,rb in zip(a.table.rows,b.table.rows):
                for ca,cb in zip(ra.cells,rb.cells):
                    if ca.text!=cb.text or ca.vertical_anchor!=cb.vertical_anchor:issues.append('comparison cell/axis changed')
        # These fragments contain connector orientation/arrowheads and boundary style.
        for path in ('./a:xfrm','./a:ln','./a:solidFill','./a:noFill'):
            aa=a._element.spPr.find(path,a._element.nsmap) if hasattr(a._element,'spPr') else None
            bb=b._element.spPr.find(path,b._element.nsmap) if hasattr(b._element,'spPr') else None
            from lxml import etree
            encode=lambda x:etree.tostring(x,method='c14n') if x is not None else None
            if encode(aa)!=encode(bb):issues.append(f'relation direction/geometry/style changed: {name}');break
    return issues
