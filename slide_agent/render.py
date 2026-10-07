import json
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_DATA_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_MARK
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_AUTO_SIZE, MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from .io import asset_path
from .catalog import ROOT, is_catalog, LICENSE_NOTICE, INVENTORY, all_registry as registry
from .template_engine import instantiate
from .validate import BODY_H, BODY_Y, CONTENT_W, HEIGHT, MARGIN, WIDTH, walk_evidence


def rgb(value):
    return RGBColor.from_string(value)


def style_font(font, theme, size, bold=False, color=None):
    font.name = theme.font_family
    font.size = Pt(size)
    font.bold = bold
    font.color.rgb = rgb(color or theme.foreground)
    rpr = font._rPr
    rpr.set("lang", "ja-JP")
    for tag in ("a:ea", "a:cs"):
        node = rpr.find(tag, rpr.nsmap)
        if node is None:
            node = OxmlElement(tag)
            rpr.append(node)
        node.set("typeface", theme.font_family)


def format_frame(tf, text, theme, size, bold=False, color=None, align=PP_ALIGN.LEFT):
    tf.clear()
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = p.space_after = Pt(0)
        p.line_spacing = 1.15
        run = p.add_run()
        run.text = line
        style_font(run.font, theme, size, bold, color)
        style_font(p.font, theme, size, bold, color)


def text_box(slide, text, x, y, w, h, theme, size, role="body", **kwargs):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = role
    format_frame(shape.text_frame, text, theme, size, **kwargs)
    return shape


def panel(slide, x, y, w, h, color, role="decoration", shape_type=MSO_AUTO_SHAPE_TYPE.RECTANGLE):
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = role
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.line.fill.background()
    # Override the default Office theme's inherited shadow effects.
    shape._element.spPr.append(OxmlElement("a:effectLst"))
    return shape


def items(slide, values, x, y, w, h, theme, color=None):
    if not values:
        return
    row_h = h / len(values)
    for index, item in enumerate(values):
        panel(slide, x, y + index * row_h + 0.12, 0.07, 0.20, theme.accent)
        text_box(slide, item.text, x + 0.23, y + index * row_h, w - 0.23, row_h - 0.12,
                 theme, theme.body_pt, color=color)


def picture(slide, path, x, y, w, h, mode):
    with Image.open(path) as im:
        iw, ih = im.size
    ratio, box_ratio = iw / ih, w / h
    if mode == "fit":
        pw, ph = (w, w / ratio) if ratio >= box_ratio else (h * ratio, h)
        return slide.shapes.add_picture(str(path), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    if ratio > box_ratio:
        pic.crop_left = pic.crop_right = (1 - box_ratio / ratio) / 2
    else:
        pic.crop_top = pic.crop_bottom = (1 - ratio / box_ratio) / 2
    return pic


def render(plan, source, base, theme, destination):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(WIDTH), Inches(HEIGHT)
    props = prs.core_properties
    props.title, props.subject = plan.title, "Source-traceable editable slides"
    props.author, props.last_modified_by = "slide-layout-agent", "slide-layout-agent"
    props.comments = "Source references and any third-party template license are recorded in slide notes."
    assets = {a.id: a for a in source.images}
    segments = {s.id: s for s in source.segments}
    for index, design in enumerate(plan.slides):
        if is_catalog(design.layout_id):
            slide = instantiate(prs, design, plan, source, index + 1, base)
            entry = registry()[design.layout_id]
            refs = sorted({r.source_id for item in walk_evidence(design) for r in item.refs})
            slide.notes_slide.notes_text_frame.text = json.dumps({
                "slide_id": design.id, "origin": design.origin, "layout_id": design.layout_id,
                "variant": design.variant, "rationale": design.rationale,
                "source_sha256": plan.source_sha256,
                "evidence": [item.model_dump(mode="json") for item in walk_evidence(design)],
                "citations": {sid: segments[sid].citation for sid in refs},
                "template_license": (ROOT/'LICENSE').read_text(encoding='utf-8') if entry.get('catalog')=='editorial' else LICENSE_NOTICE,
                "template_upstream_commit": None if entry.get('catalog')=='editorial' else INVENTORY["commit"],
                "catalog": entry.get('catalog', 'reference'),
                "family": entry.get('family'), "structural_variant": entry.get('structural_variant'),
                "repetition_reason": getattr(design, 'repetition_reason', None),
                "row_alignments": getattr(design, 'row_alignments', {}),
                "images": {key: value.model_dump() for key, value in getattr(design.contents, 'images', {}).items()},
                "template_sha256": registry()[design.layout_id]["template_sha256"],
            }, ensure_ascii=False, indent=2)
            continue
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb(theme.background)
        dark = design.layout_id in ("title", "closing")
        if dark:
            slide.background.fill.fore_color.rgb = rgb(theme.foreground)
        fg = theme.surface if dark else theme.foreground
        panel(slide, MARGIN, 0.32, 0.5, 0.055, theme.accent)
        text_box(slide, design.title.text, MARGIN, 0.50, CONTENT_W, 1.0, theme, theme.title_pt, role="title", bold=True, color=fg)
        if design.lead:
            text_box(slide, design.lead.text, MARGIN, 1.55, CONTENT_W, 0.45, theme, 18, role="lead", color=fg)
        c, kind = design.contents, design.layout_id
        if kind in ("title", "closing"):
            items(slide, c.items, MARGIN + 0.3, 2.5, 10.9, 3.5, theme, color=fg)
            panel(slide, 12.15, 2.5, 0.15, 3.5, theme.accent)
        elif kind == "bullets":
            items(slide, c.items, MARGIN + 0.2, BODY_Y, CONTENT_W - 0.4, BODY_H, theme)
        elif kind == "text_image":
            items(slide, c.items, MARGIN, BODY_Y, 5.85, BODY_H, theme)
            panel(slide, 6.85, BODY_Y, 5.88, BODY_H, theme.surface)
            pic = picture(slide, asset_path(assets[c.image_id], base), 6.97, BODY_Y + 0.12, 5.64, BODY_H - 0.24, c.image_mode)
            pic.name = f"image:{c.image_id}:{c.image_mode}"
        elif kind == "comparison":
            for x, heading, values, color in ((MARGIN, c.left_title, c.left, theme.accent), (6.82, c.right_title, c.right, theme.secondary)):
                panel(slide, x, BODY_Y, 5.91, BODY_H, theme.surface)
                panel(slide, x, BODY_Y, 5.91, 0.075, color)
                text_box(slide, heading.text, x + 0.23, BODY_Y + 0.22, 5.35, 0.65, theme, theme.body_pt, bold=True, color=color)
                items(slide, values, x + 0.23, BODY_Y + 1.1, 5.34, 3.2, theme)
        elif kind == "process":
            n = len(c.steps)
            gap, w = 0.35, (CONTENT_W - 0.35 * (n - 1)) / n
            for i, step in enumerate(c.steps):
                x = MARGIN + i * (w + gap)
                panel(slide, x, 2.65, w, 3.2, theme.surface)
                text_box(slide, str(i + 1).zfill(2), x + 0.18, 2.85, w - 0.36, 0.5, theme, 22, role="label", bold=True, color=theme.accent)
                text_box(slide, step.text, x + 0.18, 3.5, w - 0.36, 2.1, theme, theme.body_pt)
                if i < n - 1:
                    panel(slide, x + w + 0.045, 4.0, 0.25, 0.22, theme.accent, shape_type=MSO_AUTO_SHAPE_TYPE.CHEVRON)
        elif kind == "table":
            table_shape = slide.shapes.add_table(len(c.rows) + 1, len(c.headers), Inches(MARGIN), Inches(BODY_Y), Inches(CONTENT_W), Inches(BODY_H))
            table_shape.name = "native_table"
            table = table_shape.table
            for r, row in enumerate([c.headers] + c.rows):
                for col, value in enumerate(row):
                    cell = table.cell(r, col)
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = rgb(theme.foreground if r == 0 else (theme.surface if r % 2 else theme.surface_alt))
                    cell.margin_left = cell.margin_right = Inches(0.10)
                    cell.margin_top = cell.margin_bottom = Inches(0.04)
                    format_frame(cell.text_frame, value.text, theme, theme.table_pt, bold=r == 0, color=theme.surface if r == 0 else theme.foreground)
        elif kind == "chart":
            data = CategoryChartData()
            data.categories = [t.text for t in c.categories]
            for s in c.series:
                data.add_series(s.name.text, [v.value for v in s.values])
            single_series = len(c.series) == 1
            if single_series:
                text_box(slide, c.series[0].name.text, MARGIN, BODY_Y, CONTENT_W, 0.4, theme, 18, role="label", bold=True)
            chart_y = BODY_Y + (0.42 if single_series else 0)
            chart_h = 3.7 - (0.42 if single_series else 0)
            chart_shape = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(MARGIN), Inches(chart_y), Inches(CONTENT_W), Inches(chart_h), data)
            chart_shape.name = "native_chart"
            chart = chart_shape.chart
            chart.has_title = False
            chart.has_legend = not single_series
            if chart.has_legend:
                chart.legend.position = XL_LEGEND_POSITION.BOTTOM
                chart.legend.include_in_layout = False
                style_font(chart.legend.font, theme, 14)
            chart.value_axis.minimum_scale = 0
            chart.value_axis.tick_labels.number_format = "General"
            for axis in (chart.category_axis, chart.value_axis):
                style_font(axis.tick_labels.font, theme, 14)
                axis.major_tick_mark = XL_TICK_MARK.NONE
                axis.minor_tick_mark = XL_TICK_MARK.NONE
            plot = chart.plots[0]
            plot.vary_by_categories = False
            plot.has_data_labels = True
            plot.data_labels.position = XL_DATA_LABEL_POSITION.OUTSIDE_END
            plot.data_labels.number_format = "General"
            style_font(plot.data_labels.font, theme, 14)
            for i, s in enumerate(chart.series):
                s.format.fill.solid()
                s.format.fill.fore_color.rgb = rgb(theme.accent if i == 0 else theme.secondary)
                s.format.line.fill.background()
            text_box(slide, c.note.text, MARGIN, 5.85, CONTENT_W, 0.8, theme, 18, role="lead")
        refs = sorted({r.source_id for item in walk_evidence(design) for r in item.refs})
        footer = "SOURCE  " + ", ".join(refs)
        if len(footer) > 100:
            footer = f"SOURCE  {len(refs)} references — see notes"
        text_box(slide, footer, MARGIN, 6.98, 11.1, 0.3, theme, theme.footnote_pt, role="footer", color=theme.surface if dark else theme.muted)
        text_box(slide, f"{index + 1:02d}", 12.05, 6.94, 0.6, 0.35, theme, 14, role="footer", color=fg, align=PP_ALIGN.RIGHT)
        note = {
            "slide_id": design.id, "origin": design.origin, "layout_id": kind,
            "rationale": design.rationale, "source_sha256": plan.source_sha256,
            "evidence": [item.model_dump(mode="json") for item in walk_evidence(design)],
            "citations": {sid: segments[sid].citation for sid in refs},
        }
        slide.notes_slide.notes_text_frame.text = json.dumps(note, ensure_ascii=False, indent=2)
    Path(destination).parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(destination))
