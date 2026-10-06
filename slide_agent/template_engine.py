"""Populate reviewed PPTX shapes; no layout recreation at generation time.

OPC cloning is deliberately limited to native chart + embedded workbook parts.
The pinned templates have no pictures, hyperlinks, OLE, external data or macros.
Every slide gets private chart/workbook parts so repeated layouts cannot alias.
"""
import copy
import json
import math
import re
from pathlib import PurePosixPath

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.opc.package import PartFactory
from pptx.opc.packuri import PackURI
from pptx.oxml.ns import qn

from .catalog import registry, template_path

RNS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
COOL = {"322014": "071B2C", "5A3921": "1E5F8C", "C5A681": "79C8DC", "8A7B6B": "666B70",
        "E2DCD2": "D9DCDF", "EFEEE8": "E7F4F8", "A22727": "D94C68", "F2EFE9": "F2F2F0"}
ALLOWED_PARTS = {"application/vnd.openxmlformats-officedocument.drawingml.chart+xml",
                 "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
FONT_PROFILES={"source":("Yu Gothic","Yu Mincho"),"meiryo":("Meiryo","Yu Mincho"),
               "noto":("Noto Sans CJK JP","Noto Serif CJK JP"),"hiragino":("Hiragino Sans","Hiragino Mincho ProN")}


def clone_slide(prs, source):
    """Clone native shapes and their complete editable data, remapping every rId."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    package = prs.part.package
    reserved = {str(p.partname) for p in package.iter_parts()}
    memo = {}

    def clone_part(part):
        if part in memo:
            return memo[part]
        if part.content_type not in ALLOWED_PARTS:
            raise ValueError(f"unsupported template relationship part: {part.content_type}")
        old = PurePosixPath(str(part.partname))
        stem = re.sub(r"\d+$", "", old.stem)
        i = 1
        while str(old.with_name(f"{stem}{i}{old.suffix}")) in reserved:
            i += 1
        name = str(old.with_name(f"{stem}{i}{old.suffix}"))
        reserved.add(name)
        new = PartFactory(PackURI(name), part.content_type, package, part.blob)
        memo[part] = new
        mapping = {}
        for rel in part.rels.values():
            if rel.is_external:
                raise ValueError("external template relationship rejected")
            mapping[rel.rId] = new.relate_to(clone_part(rel.target_part), rel.reltype)
        if hasattr(new, "_element"):
            remap(new._element, mapping)
        elif mapping:
            raise ValueError("binary part with relationships is unsupported")
        return new

    mapping = {}
    for rel in source.part.rels.values():
        if rel.reltype.endswith(("/slideLayout", "/notesSlide")):
            continue
        if rel.is_external:
            raise ValueError("external template relationship rejected")
        mapping[rel.rId] = slide.part.relate_to(clone_part(rel.target_part), rel.reltype)
    for shape in source.shapes:
        element = copy.deepcopy(shape.element)
        remap(element, mapping)
        slide.shapes._spTree.insert_element_before(element, "p:extLst")
    if source._element.cSld.bg is not None:
        slide._element.cSld.insert(0, copy.deepcopy(source._element.cSld.bg))
    return slide


def remap(element, mapping):
    for node in element.iter():
        for key, value in list(node.attrib.items()):
            if key.startswith("{" + RNS + "}"):
                if value not in mapping:
                    raise ValueError(f"unresolved relationship {value}")
                node.set(key, mapping[value])


def replace_text(tf, text):
    """Retain the first paragraph/run style, margins, vertical alignment and bounds."""
    originals=[copy.deepcopy(p._p) for p in tf.paragraphs]
    for old in list(tf._txBody):
        if old.tag == qn("a:p"):
            tf._txBody.remove(old)
    for index,line in enumerate(text.split("\n")):
        para = copy.deepcopy(originals[min(index,len(originals)-1)])
        oldruns=para.findall(qn("a:r"))
        rpr=oldruns[0].find(qn("a:rPr")) if oldruns else None
        style=copy.deepcopy(rpr) if rpr is not None else None
        seen_ppr=False
        for child in list(para):
            if child.tag==qn("a:pPr"):
                if seen_ppr: para.remove(child)
                seen_ppr=True
            elif child.tag!=qn("a:endParaRPr"): para.remove(child)
        run = etree.Element(qn("a:r"))
        if style is not None:
            run.append(copy.deepcopy(style))
        etree.SubElement(run, qn("a:t")).text = line
        # endParaRPr is last in a:p.
        end = para.find(qn("a:endParaRPr"))
        para.insert(list(para).index(end) if end is not None else len(para), run)
        tf._txBody.append(para)


def shape_map(slide):
    return {s.name: s for s in slide.shapes}


def fill_chart(shape, content, spec):
    if spec["kind"] == "xy":
        data = XyChartData()
        for series in content.series:
            target = data.add_series(series.name.text)
            for x, y in zip(series.x_values, series.values):
                target.add_data_point(x.value, y.value)
    else:
        data = CategoryChartData()
        data.categories = [x.text for x in content.categories]
        for series in content.series:
            data.add_series(series.name.text, [x.value for x in series.values])
    shape.chart.replace_data(data)
    if spec["kind"]=="xy":
        labels=shape.chart._chartSpace.xpath(".//c:dLbl/c:tx/c:rich//a:t")
        for target,value in zip(labels,content.categories): target.text=value.text


def alias_value(spec,contents):
    chart=contents.charts[spec["chart"]]
    series=chart.series[spec["series"]]
    values=[v.value for v in series.values]
    kind=spec["kind"]
    if kind=="series_name": return series.name.text
    if kind=="category": return chart.categories[spec["point"]].text
    if kind=="x_mean": return f"平均 {sum(x.value for x in chart.series[spec['series']].x_values)/len(chart.series[spec['series']].x_values):.1f}"
    if kind=="last": return f"{values[-1]:g}"
    if kind=="sum": return f"{sum(values):+g}"
    if kind=="cagr": return f"{((values[-1]/values[0])**(1/(len(values)-1))-1)*100:.0f}%"
    raise ValueError(f"unsupported alias {kind}")


def apply_metrics(slide, entry, contents):
    shapes = shape_map(slide)
    values = {k: v.value for k, v in contents.metrics.items()}
    for key, spec in entry["metrics"].items():
        value = values[key]
        if spec.get("label"):
            replace_text(shapes[spec["label"]].text_frame, f"{value:g}{spec.get('suffix','')}")
        kind = spec["binding"]
        if kind == "dots":
            for i, name in enumerate(spec["shapes"]):
                s = shapes[name]
                s.fill.solid()
                s.fill.fore_color.rgb = RGBColor.from_string("5A3921" if i < round(value) else "E2DCD2")
        elif kind == "area":
            s = shapes[spec["shape"]]
            diameter = spec["max_diameter"] * math.sqrt(value / spec["max"])
            s.width = s.height = max(1, round(diameter))
            s.left = round(spec["center"][0] - diameter / 2)
            s.top = round(spec["center"][1] - diameter / 2)
            if diameter<210000 and spec.get("label"):
                for p in shapes[spec["label"]].text_frame.paragraphs:
                    for run in p.runs: run.font.color.rgb=RGBColor.from_string("322014")
        elif kind == "gantt_start":
            s = shapes[spec["shape"]]
            end = values[spec["end_key"]]
            s.left = round(spec["left"] + value * spec["step"])
            s.width = round((end - value) * spec["step"] - spec["gap"])
    bridge = entry.get("bridge")
    if bridge:
        vals = [values[k] for k in bridge["keys"]]
        total = vals[0]
        heights, bottoms = [total], [0]
        for delta in vals[1:-1]:
            heights.append(abs(delta))
            bottoms.append(min(total, total + delta) if bridge["cumulative"] else 0)
            total += delta
        heights.append(vals[-1]); bottoms.append(0)
        scale = bridge["height"] / max(b + h for b, h in zip(bottoms, heights))
        for name, label, h, bottom in zip(bridge["shapes"], bridge["labels"], heights, bottoms):
            s = shapes[name]
            s.top = round(bridge["baseline"] - (bottom + h) * scale)
            s.height = max(1, round(h * scale))
            shapes[label].top = s.top - bridge["label_gap"]
    calculation=entry.get("calculation")
    if calculation:
        for group in calculation["groups"]:
            maximum=max(values[b["key"]]+values.get(b.get("extra_key"),0) for b in group)
            scale=calculation["height"]/maximum
            for bar in group:
                s=shapes[bar["shape"]]; label=shapes[bar["label"]]
                s.height=max(1,round(values[bar["key"]]*scale)); s.top=calculation["baseline"]-s.height
                label.top=min(s.top,calculation["baseline"]-label.height)
                if "extra_shape" in bar:
                    extra=shapes[bar["extra_shape"]]
                    extra.height=max(1,round(values[bar["extra_key"]]*scale)); extra.top=s.top-extra.height
        for name,bar in zip(calculation["totals"],calculation["groups"][2]):
            label=shapes[name];label.top=shapes[bar["extra_shape"]].top-label.height-40000
            replace_text(label.text_frame,f"{values[bar['key']]+values[bar['extra_key']]:g}")


def apply_states(slide,entry,contents):
    shapes=shape_map(slide)
    for key,value in contents.states.items():
        spec=entry["states"][key]; s=shapes[spec["shape"]]
        color=spec["palette"][value.text]
        if spec["binding"]=="harvey":
            s.fill.solid();s.fill.fore_color.rgb=RGBColor.from_string("322014" if value.text=="満たす" else "FFFFFF")
            half=shapes[spec["half_shape"]]
            if value.text=="一部満たす":
                half.fill.solid();half.fill.fore_color.rgb=RGBColor.from_string("322014")
            else: half.fill.background()
            half.line.fill.background()
        else:
            s.fill.solid();s.fill.fore_color.rgb=RGBColor.from_string(color)
            if spec.get("label"):
                label=shapes[spec["label"]]
                replace_text(label.text_frame,value.text)
                for p in label.text_frame.paragraphs:
                    for run in p.runs: run.font.color.rgb=RGBColor.from_string("FFFFFF" if value.text=="高" else "322014")


def instantiate(prs, design, plan, source, page):
    entry = registry()[design.layout_id]
    template = Presentation(str(template_path(entry)))
    slide = clone_slide(prs, template.slides[0])
    shapes = shape_map(slide)
    replace_text(shapes["slot:title"].text_frame, design.title.text)
    for key, value in design.contents.texts.items():
        spec = entry["texts"][key]
        shape = shapes[spec["shape"]]
        tf = shape.table.cell(*spec["cell"]).text_frame if "cell" in spec else shape.text_frame
        replace_text(tf, value.text)
    for key, value in design.contents.charts.items():
        fill_chart(shapes[entry["charts"][key]["shape"]], value, entry["charts"][key])
    for name,spec in entry.get("chart_aliases",{}).items():
        replace_text(shapes[name].text_frame,alias_value(spec,design.contents))
    if entry.get("reference_line"):
        spec=entry["reference_line"]
        xs=design.contents.charts[spec["chart"]].series[0].x_values
        mean=sum(x.value for x in xs)/len(xs)
        position=round(spec["left"]+spec["width"]*(mean-spec["domain"][0])/(spec["domain"][1]-spec["domain"][0]))
        shapes[spec["shape"]].left=position
        label=shapes[spec["label"]]
        label.left=min(position+spec["label_offset"],spec["left"]+spec["width"]-label.width)
    apply_metrics(slide, entry, design.contents)
    apply_states(slide,entry,design.contents)
    for s in slide.shapes:
        if s.name == "meta:header": replace_text(s.text_frame, plan.title[:40])
        if s.name == "meta:layout": replace_text(s.text_frame, entry["part_id"])
        if s.name == "meta:footer": replace_text(s.text_frame, "slide-layout-agent")
        if s.name == "meta:page": replace_text(s.text_frame, str(page))
        if s.name == "meta:format": replace_text(s.text_frame, "native editable")
    if design.variant == "cool":
        elements = [slide._element] + [s.chart._chartSpace for s in slide.shapes if s.has_chart]
        for element in elements:
            for node in element.iter(qn("a:srgbClr")):
                node.set("val", COOL.get(node.get("val"), node.get("val")))
    if design.font_profile != "source":
        sans,serif=FONT_PROFILES[design.font_profile]
        for element in [slide._element]+[s.chart._chartSpace for s in slide.shapes if s.has_chart]:
            for node in element.iter():
                if node.tag in (qn("a:latin"),qn("a:ea"),qn("a:cs")):
                    old=node.get("typeface","")
                    node.set("typeface",serif if re.search("Mincho|Serif|Times|Georgia",old,re.I) else sans)
    return slide
