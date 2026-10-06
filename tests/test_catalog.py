import copy
import hashlib
import json
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

from pptx import Presentation
from pydantic import ValidationError

from scripts.make_catalog_samples import sample_plan
from slide_agent.catalog import INVENTORY, ROOT, registry, template_path
from slide_agent.catalog_validate import validate_catalog
from slide_agent.models import Plan, Theme
from slide_agent.render import render
from slide_agent.audit import audit
from slide_agent.cli import split_plan
from slide_agent.validate import validate


class CatalogTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.entries=list(registry().values())
        cls.source,cls.plan=sample_plan(cls.entries)

    def codes(self,slide):
        errors=[]
        validate_catalog(slide,lambda c,m,s:errors.append(c))
        return errors

    def test_exact_inventory_no_similar_layout_exclusions(self):
        self.assertEqual(len(self.entries),62)
        self.assertEqual([e["source_section"] for e in self.entries[:27]],list(range(1,28)))
        self.assertEqual([e["source_section"] for e in self.entries[27:]],list(range(1,36)))
        self.assertEqual(len({e["source_pptx_slide"] for e in self.entries}),62)
        self.assertEqual(len(INVENTORY["non_layouts"]["pptx_catalog_navigation_slides"]),8)
        self.assertIn("research_basis",registry()); self.assertIn("evidence_basis",registry())
        for source in INVENTORY['source_html_files']:
            self.assertEqual(hashlib.sha256((ROOT/'catalog/upstream'/source['path']).read_bytes()).hexdigest(),source['sha256'])

    def test_every_template_hash_native_structure_and_slots(self):
        for e in self.entries:
            with self.subTest(layout=e["layout_id"]):
                prs=Presentation(str(template_path(e)))
                self.assertEqual(len(prs.slides),1)
                s=prs.slides[0]
                self.assertEqual(len(s.shapes),e["shape_count"])
                self.assertFalse(any(x.shape_type==13 for x in s.shapes))
                names=[x.name for x in s.shapes]
                self.assertEqual(len(names),len(set(names)))
                self.assertIn("slot:title",names)
                self.assertIn("Permission is hereby granted",s.notes_slide.notes_text_frame.text)

    def test_every_layout_both_variants_generate_without_office(self):
        with tempfile.TemporaryDirectory() as tmp:
            for variant in ["warm","cool"]:
                source,plan=sample_plan(self.entries,variant)
                self.assertTrue(validate(plan,source,Path(tmp),Theme())["ok"])
                path=Path(tmp)/f"{variant}.pptx"
                render(plan,source,Path(tmp),Theme(),path)
                result=audit(path,Theme())
                self.assertTrue(result["ok"],result["issues"])
                self.assertEqual(result["counts"]["slides"],62)
                self.assertEqual(result["counts"]["pictures"],0)
                self.assertEqual(result["counts"]["charts"],10)
                self.assertEqual(result["counts"]["embedded_workbooks"],10)
                self.assertIn("Permission is hereby granted",Presentation(path).slides[0].notes_slide.notes_text_frame.text)

    def test_each_layout_rejects_overflow_missing_unknown_and_geometry(self):
        for original in self.plan.slides:
            with self.subTest(layout=original.layout_id):
                slide=original.model_copy(deep=True)
                self.assertEqual(self.codes(slide),[])
                slide.title.text="長"*1000
                self.assertIn("TEMPLATE_OVERFLOW",self.codes(slide))
                slide=original.model_copy(deep=True)
                slot=next(iter(slide.contents.texts),None)
                if slot:
                    del slide.contents.texts[slot]
                    self.assertIn("TEMPLATE_SLOTS",self.codes(slide))
                slide=original.model_copy(deep=True)
                slide.contents.texts["intruder"]=slide.title
                self.assertIn("TEMPLATE_SLOTS",self.codes(slide))
                raw=original.model_dump(mode="json"); raw["contents"]["x"]=0
                data=self.plan.model_dump(mode="json"); data["slides"]=[raw]
                with self.assertRaises(ValidationError): Plan.model_validate(data)

    def test_each_metric_range_and_chart_cardinality(self):
        for original in self.plan.slides:
            for key in original.contents.metrics:
                with self.subTest(layout=original.layout_id,metric=key):
                    slide=original.model_copy(deep=True)
                    slide.contents.metrics[key].value=1e10
                    self.assertIn("TEMPLATE_METRIC_RANGE",self.codes(slide))
            for key in original.contents.charts:
                with self.subTest(layout=original.layout_id,chart=key):
                    slide=original.model_copy(deep=True)
                    slide.contents.charts[key].series[0].values.pop()
                    self.assertIn("TEMPLATE_CHART_DIMENSIONS",self.codes(slide))

    def test_repeated_template_charts_have_independent_workbooks(self):
        entry=registry()["chart_insight"]
        source,plan=sample_plan([entry,entry])
        key=next(iter(plan.slides[1].contents.charts))
        plan.slides[1].contents.charts[key].series[0].values[0].value=17.0
        with tempfile.TemporaryDirectory() as tmp:
            dest=Path(tmp)/"repeat.pptx"
            render(plan,source,Path(tmp),Theme(),dest)
            p=Presentation(dest)
            charts=[next(s.chart for s in sl.shapes if s.has_chart) for sl in p.slides]
            self.assertEqual(charts[0].series[0].values[0],32)
            self.assertEqual(charts[1].series[0].values[0],17)
            self.assertNotEqual(charts[0].part.partname,charts[1].part.partname)
            with ZipFile(dest) as z:
                self.assertEqual(len([n for n in z.namelist() if n.startswith("ppt/embeddings/")]),2)
                self.assertEqual(len(z.namelist()),len(set(z.namelist())))

    def test_fixed_geometry_audit_detects_mutation(self):
        source,plan=sample_plan([registry()["issue_tree"]])
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/"mutated.pptx"; render(plan,source,Path(tmp),Theme(),path)
            p=Presentation(path); p.slides[0].shapes[0].left+=10000; p.save(path)
            self.assertFalse(audit(path,Theme())["ok"])

    def test_catalog_split_rejects_overflow_with_actionable_message(self):
        p=self.plan.model_copy(deep=True); p.slides[0].title.text="長"*1000
        with self.assertRaisesRegex(ValueError,"explicitly replan"): split_plan(p,Theme())
        p=self.plan.model_copy(deep=True);p.title='長'*40
        self.assertIn('CATALOG_HEADER_OVERFLOW',{x['code'] for x in validate(p,self.source,ROOT,Theme())['issues']})

    def test_chart_reference_evidence_walked_and_numbers_checked(self):
        p=self.plan.model_copy(deep=True)
        s=next(s for s in p.slides if s.layout_id=="chart_insight")
        next(iter(s.contents.charts.values())).series[0].values[0].value=123456.0
        codes={x["code"] for x in validate(p,self.source,ROOT,Theme())["issues"]}
        self.assertIn("UNSUPPORTED_NUMBER",codes)

    def test_no_office_or_com_imports_in_runtime(self):
        for file in (ROOT/"slide_agent").glob("*.py"):
            text=file.read_text(encoding="utf-8")
            self.assertNotIn("import win32com",text)
            self.assertNotIn("import comtypes",text)

    def test_all_state_slots_reject_unknown_state(self):
        for original in self.plan.slides:
            for key in original.contents.states:
                with self.subTest(layout=original.layout_id,state=key):
                    s=original.model_copy(deep=True); s.contents.states[key].text="unknown"
                    self.assertIn("TEMPLATE_STATE",self.codes(s))

    def test_numeric_semantics_reject_inconsistent_data(self):
        cases=[("ranked_bar_annotated","TEMPLATE_RANK_ORDER"),
               ("scenario_lines_cagr","TEMPLATE_CAGR_DOMAIN"),
               ("scatter_annotated","TEMPLATE_CHART_X_DOMAIN"),
               ("true_waterfall","TEMPLATE_BRIDGE_TOTAL"),
               ("calc_flow","TEMPLATE_CALCULATION"),
               ("gantt","TEMPLATE_GANTT_ORDER")]
        for layout,code in cases:
            with self.subTest(layout=layout):
                s=next(x for x in self.plan.slides if x.layout_id==layout).model_copy(deep=True)
                if layout=="ranked_bar_annotated": next(iter(s.contents.charts.values())).series[0].values[0].value=0
                elif layout=="scenario_lines_cagr": next(iter(s.contents.charts.values())).series[0].values[0].value=0
                elif layout=="scatter_annotated": next(iter(s.contents.charts.values())).series[0].x_values[0].value=101
                elif layout=="true_waterfall": s.contents.metrics[registry()[layout]["bridge"]["keys"][-1]].value+=1
                elif layout=="calc_flow": s.contents.metrics["increment_1"].value+=1
                else: s.contents.metrics["task_1_end"].value=0
                self.assertIn(code,self.codes(s))

    def test_area_dot_and_scatter_boundary_geometry(self):
        for layout in ["proportional_circles","progress_bubble_matrix","dot_matrix_share","scatter_annotated"]:
            for boundary in ["min","max"]:
                with self.subTest(layout=layout,boundary=boundary),tempfile.TemporaryDirectory() as tmp:
                    e=registry()[layout]; source,plan=sample_plan([e]); s=plan.slides[0]
                    if layout=="scatter_annotated":
                        for v in next(iter(s.contents.charts.values())).series[0].x_values: v.value=0 if boundary=="min" else 100
                    else:
                        for k,v in s.contents.metrics.items(): v.value=float(e["metrics"][k][boundary])
                    self.assertEqual(self.codes(s),[])
                    dest=Path(tmp)/"edge.pptx";render(plan,source,Path(tmp),Theme(),dest)
                    result=audit(dest,Theme()); self.assertTrue(result["ok"],result["issues"])

    def test_manifest_documents_every_html_comparison(self):
        comparisons=json.loads((ROOT/"catalog/html-comparison.json").read_text(encoding="utf-8"))
        self.assertEqual(set(comparisons),set(registry()))
        for e in self.entries: self.assertTrue(set(comparisons[e["layout_id"]])<=set(e["differences"]))

    def test_all_layouts_replace_sample_text_numbers_and_chart_data(self):
        entries=copy.deepcopy(self.entries)
        for e in entries:
            e['title']['sample']='差込検証'
            for slot in e['texts'].values(): slot['sample']='改'
            for spec in e['charts'].values():
                spec['sample']['categories']=['甲']*spec['category_count']
                for series in spec['sample']['series']:
                    series['name']='乙';series['values']=[10.0]*len(series['values'])
                    series['x_values']=[50.0]*len(series['x_values'])
            for spec in e['metrics'].values():
                if spec['binding'] in ('dots','area'): spec['sample']=float(spec['max'])/2
            if e.get('bridge'):
                for i,k in enumerate(e['bridge']['keys']): e['metrics'][k]['sample']=100.0 if i in (0,len(e['bridge']['keys'])-1) else 0.0
            if e.get('calculation'):
                for a,b,result in zip(*e['calculation']['groups']):
                    e['metrics'][a['key']]['sample']=2.0;e['metrics'][b['key']]['sample']=3.0
                    e['metrics'][result['key']]['sample']=4.0;e['metrics'][result['extra_key']]['sample']=2.0
        source,plan=sample_plan(entries)
        with tempfile.TemporaryDirectory() as tmp:
            self.assertTrue(validate(plan,source,Path(tmp),Theme())['ok'])
            dest=Path(tmp)/'substituted.pptx';render(plan,source,Path(tmp),Theme(),dest)
            self.assertTrue(audit(dest,Theme())['ok'])
            for entry,design,slide in zip(entries,plan.slides,Presentation(dest).slides):
                shapes={s.name:s for s in slide.shapes}
                self.assertEqual(shapes['slot:title'].text,design.title.text)
                for k,spec in entry['texts'].items():
                    sh=shapes[spec['shape']]
                    actual=sh.table.cell(*spec['cell']).text if 'cell' in spec else sh.text
                    self.assertEqual(actual,design.contents.texts[k].text,(entry['layout_id'],k))
                for k,spec in entry['metrics'].items():
                    if spec.get('label'):
                        value=design.contents.metrics[k].value
                        prefix='+' if spec.get('signed') and value>0 else ''
                        self.assertEqual(shapes[spec['label']].text,prefix+f"{value:g}{spec.get('suffix','')}")
                for k,spec in entry['charts'].items():
                    chart=shapes[spec['shape']].chart
                    for series in chart.series:
                        self.assertEqual(set(series.values),{10.0})
                        self.assertEqual(series.name,'乙')
                    # The private editable workbook must also be refreshed.
                    from io import BytesIO
                    with ZipFile(BytesIO(chart.part.chart_workbook.xlsx_part.blob)) as workbook:
                        shared=workbook.read('xl/sharedStrings.xml').decode('utf-8')
                        self.assertIn('乙',shared)
                        self.assertNotIn('ラベル',shared)

    def test_authored_catalog_story_uses_saved_templates(self):
        from scripts.make_catalog_story import make_story
        with tempfile.TemporaryDirectory() as tmp:
            source,plan=make_story(tmp)
            self.assertEqual([s.layout_id for s in plan.slides],['title_page','chart_insight','decision_page'])
            self.assertTrue(validate(plan,source,Path(tmp),Theme())['ok'])


if __name__=="__main__": unittest.main()
