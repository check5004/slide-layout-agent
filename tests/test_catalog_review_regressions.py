import copy
import json
import shutil
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from pptx import Presentation
from pydantic import ValidationError

from scripts.make_catalog_samples import sample_plan
from slide_agent.audit import audit
from slide_agent.catalog import ROOT,registry
from slide_agent.catalog_validate import validate_catalog
from slide_agent.io import fingerprint
from slide_agent.models import Theme,ElapsedYears
from slide_agent.numeric_format import exact_number
from slide_agent.render import render
from slide_agent.schema_contract import layout_schema
from slide_agent.template_engine import alias_value
from slide_agent.validate import validate


def review_entry(layout,values=None,years=None):
    e=copy.deepcopy(registry()[layout])
    if values is not None:
        for k,v in zip(e['bridge']['keys'],values):e['metrics'][k]['sample']=v
    if years is not None:
        sample=next(iter(e['charts'].values()))['sample']
        sample['categories']=[str(y) for y in range(2020,2031,2)]
        sample['elapsed_years']=years
        e['texts']['text_012']['sample']='単位、2020〜2030年'
    return e


class ReviewRegressions(unittest.TestCase):
    def codes(self,slide):
        errors=[];validate_catalog(slide,lambda c,m,s:errors.append(c));return errors

    def test_cagr_uses_explicit_ten_year_duration_in_both_palettes(self):
        e=review_entry('scenario_lines_cagr',years=10)
        for variant in ('warm','cool'):
            with self.subTest(variant=variant),tempfile.TemporaryDirectory() as temp:
                source,plan=sample_plan([e],variant);dest=Path(temp)/'cagr.pptx'
                self.assertTrue(validate(plan,source,Path(temp),Theme())['ok'])
                render(plan,source,Path(temp),Theme(),dest)
                self.assertTrue(audit(dest,Theme())['ok'])
                shapes={s.name:s for s in Presentation(dest).slides[0].shapes}
                self.assertEqual([shapes[f'shape:{i:03}'].text for i in (19,21,23)],['13%','11%','9%'])
                self.assertEqual(shapes['shape:017'].text,'年平均成長率（10年）')

    def test_cagr_missing_mismatched_ambiguous_period_rejected(self):
        e=review_entry('scenario_lines_cagr',years=10);_,plan=sample_plan([e]);s=plan.slides[0]
        chart=next(iter(s.contents.charts.values()))
        chart.elapsed_years.value=5
        self.assertIn('TEMPLATE_CAGR_PERIOD',self.codes(s))
        chart.elapsed_years=None
        self.assertIn('TEMPLATE_CAGR_PERIOD',self.codes(s))
        for category in chart.categories:category.text='stage'
        self.assertIn('TEMPLATE_CAGR_PERIOD',self.codes(s))
        with self.assertRaisesRegex(ValueError,'explicit'):
            alias_value(next(v for v in e['chart_aliases'].values() if v['kind']=='cagr'),s.contents)
        chart.elapsed_years=ElapsedYears(value=10.,refs=chart.series[0].values[0].refs)
        self.assertEqual(self.codes(s),[])
        for category,value in zip(chart.categories,['2020','2021','2024','2026','2028','2030']):category.text=value
        self.assertIn('TEMPLATE_CAGR_CALENDAR',self.codes(s))
        for invalid in (0.,-1.,float('nan')):
            with self.assertRaises(ValidationError):ElapsedYears(value=invalid,refs=chart.series[0].values[0].refs)

    def test_bridge_sign_controls_color_geometry_and_labels(self):
        for layout,values in [('true_waterfall',[11,3,-4,3,13]),('waterfall',[100,-18,24,9,7,122])]:
            e=review_entry(layout,values=values)
            for variant in ('warm','cool'):
                with self.subTest(layout=layout,variant=variant),tempfile.TemporaryDirectory() as temp:
                    source,plan=sample_plan([e],variant);dest=Path(temp)/'sign.pptx'
                    self.assertTrue(validate(plan,source,Path(temp),Theme())['ok'])
                    render(plan,source,Path(temp),Theme(),dest)
                    self.assertTrue(audit(dest,Theme())['ok'])
                    shapes={s.name:s for s in Presentation(dest).slides[0].shapes}
                    for i,value in enumerate(values[1:-1],1):
                        bar=shapes[e['bridge']['shapes'][i]]
                        expected=('5A3921' if value>0 else 'A22727') if variant=='warm' else ('1E5F8C' if value>0 else 'D94C68')
                        self.assertEqual(str(bar.fill.fore_color.rgb),expected)
                        self.assertEqual(shapes[e['bridge']['labels'][i]].text,('+' if value>0 else '')+str(value))
                    if layout=='true_waterfall':
                        scale=e['bridge']['height']/14;baseline=e['bridge']['baseline']
                        for i,top,height in [(1,14,3),(2,14,4),(3,13,3)]:
                            bar=shapes[e['bridge']['shapes'][i]]
                            self.assertEqual(bar.top,round(baseline-top*scale));self.assertEqual(bar.height,round(height*scale))

    def test_metric_precision_is_preserved_or_capacity_rejects(self):
        e=review_entry('true_waterfall',values=[123456.7,1,2,3,123462.7])
        for variant in ('warm','cool'):
            with self.subTest(variant=variant),tempfile.TemporaryDirectory() as temp:
                source,plan=sample_plan([e],variant);dest=Path(temp)/'precision.pptx'
                self.assertTrue(validate(plan,source,Path(temp),Theme())['ok'])
                render(plan,source,Path(temp),Theme(),dest)
                shapes={s.name:s for s in Presentation(dest).slides[0].shapes}
                self.assertEqual([shapes[n].text for n in e['bridge']['labels']],['123456.7','+1','+2','+3','123462.7'])
        self.assertEqual(exact_number(.00000123456789),'0.00000123456789')
        self.assertEqual(exact_number(123456.700),'123456.7')
        _,plan=sample_plan([registry()['progress_bubble_matrix']])
        next(iter(plan.slides[0].contents.metrics.values())).value=12.3456789012345
        self.assertIn('TEMPLATE_OVERFLOW',self.codes(plan.slides[0]))
        for layout in ('scenario_lines_cagr','scatter_annotated'):
            for spec in registry()[layout]['chart_aliases'].values():
                if spec['kind'] in ('cagr','x_mean'):self.assertEqual(spec['number_format']['rounding'],'half_even')

    def test_numeric_quote_does_not_claim_qualifier_coverage(self):
        source,plan=sample_plan([registry()['chart_insight']])
        value=next(iter(plan.slides[0].contents.charts.values())).series[0].values[0]
        qualifier='estimate only; causal effect not verified'
        quote='32.0 ('+qualifier+')'
        source.segments[0].text=source.segments[0].text.replace('32',quote,1)
        value.refs[0].quote=quote;plan.source_sha256=fingerprint(source)
        report=validate(plan,source,ROOT,Theme());codes={i['code'] for i in report['issues']}
        self.assertTrue(report['ok']);self.assertIn('PARTIAL_SOURCE_REVIEW',codes);self.assertIn('NUMERIC_CONTEXT_REVIEW',codes)
        self.assertEqual(report['mechanical_validation'],'passed');self.assertEqual(report['source_coverage'],'review_required')
        self.assertEqual(report['semantic_review'],'not_performed')
        body=plan.slides[0].contents.texts['text_026'];body.text=qualifier;body.refs=value.refs
        report=validate(plan,source,ROOT,Theme())
        self.assertNotIn('NUMERIC_CONTEXT_REVIEW',{i['code'] for i in report['issues']})
        self.assertEqual(report['semantic_review'],'not_performed')

    def test_saved_schema_and_contract_drift_is_rejected(self):
        e=registry()['issue_tree'];source,plan=sample_plan([e])
        with tempfile.TemporaryDirectory() as temp:
            base=Path(temp);(base/'catalog/schemas').mkdir(parents=True)
            shutil.copyfile(ROOT/'catalog/plan.schema.json',base/'catalog/plan.schema.json')
            path=base/'catalog/schemas/issue_tree.schema.json';data=layout_schema(e)
            data['properties']['layout_id']['const']='wrong';path.write_text(json.dumps(data),encoding='utf-8')
            with patch('slide_agent.schema_contract.ROOT',base):
                report=validate(plan,source,ROOT,Theme())
                self.assertFalse(report['ok']);self.assertIn('CATALOG_SCHEMA_MISMATCH',{i['code'] for i in report['issues']})
        changed=copy.deepcopy(e);changed['title']['max_lines']+=1
        self.assertNotEqual(layout_schema(changed)['x-contract-sha256'],layout_schema(e)['x-contract-sha256'])

    def test_catalog_cp932_stdout_and_direct_utf8_file(self):
        env={**os.environ,'PYTHONIOENCODING':'cp932','PYTHONUTF8':'0'}
        command=[sys.executable,'-B','-m','slide_agent','catalog','--layout','executive_summary']
        result=subprocess.run(command,cwd=ROOT,env=env,capture_output=True)
        self.assertEqual(result.returncode,0,result.stderr)
        self.assertTrue(all(byte<128 for byte in result.stdout))
        self.assertEqual(json.loads(result.stdout)['name'],registry()['executive_summary']['name'])
        with tempfile.TemporaryDirectory() as temp:
            path=Path(temp)/'catalog.json'
            result=subprocess.run(command+['--out',str(path)],cwd=ROOT,env=env,capture_output=True)
            self.assertEqual(result.returncode,0,result.stderr)
            raw=path.read_text(encoding='utf-8')
            self.assertIn(registry()['executive_summary']['name'],raw)
            self.assertEqual(json.loads(raw)['layout_id'],'executive_summary')
            refusal=subprocess.run(command+['--out',str(path)],cwd=ROOT,env=env,capture_output=True)
            self.assertEqual(refusal.returncode,2)


if __name__=='__main__':unittest.main()
