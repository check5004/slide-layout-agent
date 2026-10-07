import copy
import json
import math
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from pptx import Presentation
from pydantic import ValidationError

from scripts.make_editorial_samples import editorial_sample_plan
from scripts.make_editorial_story import story
from slide_agent.audit import audit
from slide_agent.catalog import ROOT, editorial_registry, registry, template_path
from slide_agent.catalog_validate import validate_catalog
from slide_agent.io import fingerprint
from slide_agent.models import Plan, Theme
from slide_agent.render import render
from slide_agent.selection import candidates, select_variants, selection_report
from slide_agent.schema_contract import layout_schema
from slide_agent.validate import validate, units


class EditorialTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.base = Path(cls.temp.name)
        cls.entries = list(editorial_registry().values())
        cls.source, cls.plan = editorial_sample_plan(cls.entries, cls.base)

    @classmethod
    def tearDownClass(cls): cls.temp.cleanup()

    def codes(self, slide):
        errors=[]; validate_catalog(slide,lambda c,m,s: errors.append(c)); return errors

    def sample(self, key):
        return editorial_sample_plan([editorial_registry()[key]], self.base)

    def test_original_catalog_is_isolated_and_meaningful(self):
        self.assertEqual(len(registry()), 62)
        self.assertEqual(len({e['family'] for e in self.entries}), 14)
        self.assertGreaterEqual(len(self.entries), 64)
        self.assertEqual(len({e['layout_id'] for e in self.entries}), len(self.entries))
        for family in ('cards', 'bullets', 'steps'):
            self.assertEqual({e['item_count'] for e in self.entries if e['family']==family}, {2,3,4,5,6})
        self.assertTrue(all(e['slot_mapping'] and e['semantic_signature'] for e in self.entries))

    def test_every_template_native_and_named_slots(self):
        for entry in self.entries:
            with self.subTest(layout=entry['layout_id']):
                prs=Presentation(template_path(entry)); slide=prs.slides[0]
                self.assertEqual(len(prs.slides),1)
                names={s.name for s in slide.shapes}
                self.assertEqual(len(names),len(slide.shapes))
                self.assertEqual(len(slide.shapes),entry['shape_count'])
                for spec in entry['texts'].values():
                    self.assertIn(spec['shape'],names)
                    self.assertGreaterEqual(spec['font_pt'],12)
                    if spec['role']=='body': self.assertGreaterEqual(spec['font_pt'],18)
                self.assertFalse(any(s.shape_type==13 for s in slide.shapes))
                self.assertIn('Original design',slide.notes_slide.notes_text_frame.text)

    def test_all_generate_without_office_or_subprocess(self):
        import builtins
        original=builtins.__import__
        def guarded(name,*args,**kwargs):
            if name.startswith(('win32com','comtypes')): raise AssertionError('Office dependency')
            return original(name,*args,**kwargs)
        with patch('builtins.__import__',side_effect=guarded), patch('subprocess.Popen',side_effect=AssertionError('external process')):
            self.assertTrue(validate(self.plan,self.source,self.base,Theme())['ok'])
            dest=self.base/'all.pptx';render(self.plan,self.source,self.base,Theme(),dest)
            result=audit(dest,Theme());self.assertTrue(result['ok'],result['issues'])
            self.assertEqual(result['counts']['slides'],len(self.entries))
            self.assertEqual(result['counts']['pictures'],2)
            self.assertEqual(result['counts']['charts'],4)
            self.assertEqual(result['counts']['tables'],4)
            self.assertEqual(result['counts']['embedded_workbooks'],4)

    def test_each_slot_exact_capacity_and_overflow_no_shrink(self):
        for original, entry in zip(self.plan.slides,self.entries):
            with self.subTest(layout=original.layout_id):
                self.assertEqual(self.codes(original),[])
                for key,spec in entry['texts'].items():
                    slide=original.model_copy(deep=True)
                    text='調'*math.floor(spec['max_units_per_line'])
                    slide.contents.texts[key].text='\n'.join([text]*spec['max_lines'])
                    self.assertNotIn('TEMPLATE_OVERFLOW',self.codes(slide))
                    slide.contents.texts[key].text+='\n調'
                    self.assertIn('TEMPLATE_OVERFLOW',self.codes(slide))
                slide=original.model_copy(deep=True);slide.title.text='長'*1000
                self.assertIn('TEMPLATE_OVERFLOW',self.codes(slide))

    def test_missing_unknown_empty_and_item_count_are_rejected(self):
        for original in self.plan.slides:
            with self.subTest(layout=original.layout_id):
                slide=original.model_copy(deep=True);key=next(iter(slide.contents.texts))
                del slide.contents.texts[key];self.assertIn('TEMPLATE_SLOTS',self.codes(slide))
                slide=original.model_copy(deep=True);slide.contents.texts['unknown']=slide.title
                self.assertIn('TEMPLATE_SLOTS',self.codes(slide))
                slide=original.model_copy(deep=True);slide.contents.texts[key].text='  \n '
                self.assertIn('TEMPLATE_EMPTY',self.codes(slide))
                slide=original.model_copy(deep=True);slide.title.text='  \n '
                self.assertIn('TEMPLATE_EMPTY',self.codes(slide))
        _,plan=self.sample('ed_cards_grid_3'); raw=plan.model_dump(mode='json')
        raw['slides'][0]['contents']['texts']['item_1_body']['text']=''
        with self.assertRaises(ValidationError):Plan.model_validate(raw)
        for n in (2,3,4,5,6):
            _,plan=self.sample(f'ed_cards_grid_{n}')
            fit,_=candidates(plan.slides[0]);self.assertTrue(fit)
            self.assertTrue(all(e['item_count']==n for e in fit))

    def test_boundary_fixture_preserves_all_titles_including_emphasis(self):
        from scripts.make_editorial_boundary_samples import boundary_entries
        entries=boundary_entries()
        _,plan=editorial_sample_plan(entries,self.base,boundary_fixture=True)
        checked=0
        for slide,entry in zip(plan.slides,entries):
            for key,spec in [('title',entry['title']),*entry['texts'].items()]:
                text=slide.title.text if key=='title' else slide.contents.texts[key].text
                self.assertEqual(len(text.split('\n')),spec['max_lines'])
                self.assertTrue(all(units(line)==math.floor(spec['max_units_per_line']) for line in text.split('\n')))
                checked+=1
        self.assertEqual(checked,773)

    def test_mirrors_keep_semantic_mapping_and_order(self):
        entries=editorial_registry()
        for a,b in [('ed_cards_feature_left_3','ed_cards_feature_right_3'),
                    ('ed_comparison_columns_2','ed_comparison_swapped_2'),
                    ('ed_diagram_left_3','ed_diagram_right_3'),
                    ('ed_cards_image_left_2','ed_cards_image_right_2')]:
            self.assertEqual(entries[a]['semantic_signature'],entries[b]['semantic_signature'])
            self.assertEqual(entries[a]['visual_signature'],entries[b]['visual_signature'])
            key=next(k for k in entries[a]['texts'] if k!='context')
            self.assertNotEqual(entries[a]['texts'][key]['bounds_emu'],entries[b]['texts'][key]['bounds_emu'])

    def test_selector_preserves_all_content_refs_and_avoids_four_same(self):
        source, initial, selected, decisions=story()
        self.assertEqual(len(selected.slides),16)
        for before,after in zip(initial.slides,selected.slides):
            self.assertEqual(before.contents.model_dump(),after.contents.model_dump())
            self.assertEqual(before.title,after.title);self.assertEqual(before.origin,after.origin)
            self.assertEqual(before.rationale,after.rationale)
        a,_=select_variants(initial);b,_=select_variants(initial)
        self.assertEqual(a,b)
        self.assertEqual([s.layout_id for s in initial.slides[2:6]],['ed_cards_grid_3']*4)
        signatures=[editorial_registry()[s.layout_id]['visual_signature'] for s in selected.slides[2:6]]
        self.assertGreaterEqual(len(set(signatures)),2)
        self.assertTrue(all('feature' not in s.layout_id for s in selected.slides[2:6]))
        report=validate(selected,source,ROOT/'examples/editorial-story',Theme())
        self.assertTrue(report['ok'],report['issues'])
        self.assertTrue(all(not r['fourth_consecutive_structure'] for r in report['layout_selection']))
        self.assertIn('VARIANT_REPETITION',{i['code'] for i in validate(initial,source,self.base,Theme())['issues']})

    def test_mirror_only_rotation_cannot_evade_repetition(self):
        _,plan=self.sample('ed_diagram_left_3')
        slides=[]
        for i in range(4):
            s=plan.slides[0].model_copy(deep=True);s.id=f'x{i}';s.repetition_reason=None
            s.layout_id='ed_diagram_left_3' if i%2==0 else 'ed_diagram_right_3';slides.append(s)
        plan.slides=slides;issues=[]
        selection_report(plan,lambda c,m,s,*args:issues.append(c))
        self.assertIn('VARIANT_REPETITION',issues)
        chosen,_=select_variants(plan)
        self.assertIsNotNone(chosen.slides[-1].repetition_reason)

    def test_constraints_and_unavoidable_repeat_are_visible(self):
        _,plan=self.sample('ed_cards_grid_3')
        slides=[]
        for i in range(4):
            s=plan.slides[0].model_copy(deep=True);s.id=f'x{i}';s.repetition_reason=None
            s.allowed_layouts=[s.layout_id];s.constraints_reason='Parallel record comparison requires aligned columns.'
            slides.append(s)
        plan.slides=slides
        chosen,decisions=select_variants(plan)
        self.assertTrue(chosen.slides[-1].repetition_reason)
        self.assertIn('Parallel record',chosen.slides[-1].repetition_reason)
        issues=[];report=selection_report(chosen,lambda c,m,s,*args:issues.append(c))
        self.assertIn('VARIANT_REPETITION_ACCEPTED',issues)
        self.assertTrue(report[-1]['rejected_candidates'])
        chosen.slides[0].allowed_layouts=['ed_table_two_columns_3']
        issues=[];selection_report(chosen,lambda c,m,s,*args:issues.append(c))
        self.assertIn('VARIANT_CONSTRAINT',issues)

    def test_overflow_does_not_force_another_family_or_drop_content(self):
        _,plan=self.sample('ed_cards_grid_3');slide=plan.slides[0]
        slide.contents.texts['item_1_body'].text='非常に長い日本語の調査結果と留保。'*100
        with self.assertRaisesRegex(ValueError,'no equivalent variant fits'):select_variants(plan)

    def test_parallel_items_never_gain_unsourced_hierarchy(self):
        source,plan=self.sample('ed_cards_grid_3')
        fitting,rejected=candidates(plan.slides[0])
        self.assertFalse(any('feature' in e['layout_id'] for e in fitting))
        self.assertTrue(any('hierarchy' in str(e) for e in rejected))
        plan.slides[0].layout_id='ed_cards_feature_left_3'
        self.assertIn('EDITORIAL_EMPHASIS',self.codes(plan.slides[0]))
        source,plan=self.sample('ed_cards_feature_left_3')
        self.assertNotIn('EDITORIAL_EMPHASIS',self.codes(plan.slides[0]))
        fitting,_=candidates(plan.slides[0])
        self.assertTrue(all(e['emphasis_item']=='item_1' for e in fitting))
        plan.slides[0].emphasis.refs[0].quote='Invented priority'
        self.assertIn('EMPHASIS_EVIDENCE',{i['code'] for i in validate(plan,source,self.base,Theme())['issues']})

    def test_schema_drift_includes_selection_and_image_contracts(self):
        entry=copy.deepcopy(self.entries[0]);original=layout_schema(entry)['x-contract-sha256']
        entry['visual_signature']='edited';self.assertNotEqual(original,layout_schema(entry)['x-contract-sha256'])
        entry=copy.deepcopy(editorial_registry()['ed_cards_image_left_2']);original=layout_schema(entry)['x-contract-sha256']
        entry['images']['image']['bounds_emu'][0]+=100
        self.assertNotEqual(original,layout_schema(entry)['x-contract-sha256'])

    def test_image_fit_crop_missing_and_bad_hash(self):
        source,plan=self.sample('ed_cards_image_left_2')
        for mode in ('fit','crop'):
            plan.slides[0].contents.images['image'].image_mode=mode
            dest=self.base/f'image-{mode}.pptx';render(plan,source,self.base,Theme(),dest)
            self.assertTrue(audit(dest,Theme())['ok'])
        bad=plan.model_copy(deep=True);bad.slides[0].contents.images['image'].image_id='missing'
        self.assertIn('UNKNOWN_IMAGE',{i['code'] for i in validate(bad,source,self.base,Theme())['issues']})
        del bad.slides[0].contents.images['image'];self.assertIn('TEMPLATE_IMAGE_SLOTS',self.codes(bad.slides[0]))
        bad_source=source.model_copy(deep=True);bad_source.images[0].sha256='0'*64
        self.assertIn('IMAGE_INVALID',{i['code'] for i in validate(plan,bad_source,self.base,Theme())['issues']})

    def test_repeated_charts_get_independent_workbooks(self):
        source,plan=editorial_sample_plan([editorial_registry()['ed_chart_bar_left_4']]*2,self.base)
        dest=self.base/'charts.pptx';render(plan,source,self.base,Theme(),dest)
        prs=Presentation(dest);charts=[next(s.chart for s in slide.shapes if s.has_chart) for slide in prs.slides]
        self.assertNotEqual(charts[0].part.partname,charts[1].part.partname)
        with ZipFile(dest) as z:self.assertEqual(sum(n.startswith('ppt/embeddings/') for n in z.namelist()),2)

    def test_totals_and_comparison_axes_are_semantic_constraints(self):
        _,plan=self.sample('ed_metrics_breakdown_top_3')
        self.assertEqual(self.codes(plan.slides[0]),[])
        plan.slides[0].contents.metrics['total'].value+=1
        self.assertIn('EDITORIAL_METRIC_TOTAL',self.codes(plan.slides[0]))
        _,plan=self.sample('ed_comparison_columns_2')
        plan.slides[0].contents.texts['b_axis_1'].text='Different axis'
        self.assertIn('EDITORIAL_COMPARISON_AXES',self.codes(plan.slides[0]))

    def test_text_flow_keeps_caveat_near_body_and_rejects_escape(self):
        source,plan=self.sample('ed_cards_grid_3')
        # A short source-backed fixture still preserves all exact refs.
        dest=self.base/'flow.pptx';render(plan,source,self.base,Theme(),dest)
        prs=Presentation(dest);shapes={s.name:s for s in prs.slides[0].shapes}
        body=shapes['slot:item_1_body'];tail=shapes['slot:item_1_evidence']
        self.assertLessEqual(tail.top-body.top-body.height,200000)
        self.assertTrue(audit(dest,Theme())['ok'])
        tail.top=body.top-1000;prs.save(dest)
        self.assertFalse(audit(dest,Theme())['ok'])
        tail._element.getparent().remove(tail._element);prs.save(dest)
        self.assertFalse(audit(dest,Theme())['ok'])

    def test_cli_defaults_to_new_catalog_and_keeps_reference(self):
        for args,expected in [([],len(self.entries)),(['--collection','reference'],62),(['--family','cards'],16)]:
            result=subprocess.run([sys.executable,'-m','slide_agent','catalog',*args],cwd=ROOT,capture_output=True)
            self.assertEqual(result.returncode,0,result.stderr)
            self.assertEqual(len(json.loads(result.stdout)),expected)


if __name__=='__main__':unittest.main()
