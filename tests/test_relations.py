import copy
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

from pptx import Presentation
from pptx.oxml.ns import qn
from pydantic import ValidationError
from scripts.make_relation_samples import fixture
from slide_agent.catalog import ROOT, relation_registry, template_path
from slide_agent.models import Plan, Theme
from slide_agent.catalog_validate import validate_catalog
from slide_agent.schema_contract import layout_schema
from slide_agent.selection import select_variants
from slide_agent.render import render
from slide_agent.audit import audit
from slide_agent.validate import validate


class RelationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp=tempfile.TemporaryDirectory();cls.base=Path(cls.temp.name)
        cls.entries=list(relation_registry().values())
    @classmethod
    def tearDownClass(cls):cls.temp.cleanup()
    def sample(self,key,stress=False):return fixture([relation_registry()[key]],stress)
    def codes(self,slide):
        result=[];validate_catalog(slide,lambda c,m,s:result.append(c));return result
    def deck(self,source,plan,name='result.pptx'):
        path=self.base/name;render(plan,source,self.base,Theme(),path);return path

    def test_all_25_templates_and_boundaries_generate_without_office(self):
        self.assertEqual(len(self.entries),25)
        for stress in (False,True):
            source,plan=fixture(self.entries,stress)
            report=validate(plan,source,self.base,Theme());self.assertTrue(report['ok'],report['issues'])
            path=self.deck(source,plan)
            result=audit(path,Theme());self.assertTrue(result['ok'],result['issues'])
            self.assertEqual(result['counts']['pictures'],0)
            with ZipFile(path) as z:self.assertGreater(sum(b'cxnSp' in z.read(n) for n in z.namelist() if n.startswith('ppt/slides/slide') and n.endswith('.xml')),20)
        for e in self.entries:
            self.assertEqual(json.loads((ROOT/e['schema']).read_text(encoding='utf-8')),layout_schema(e))
            self.assertEqual(len(Presentation(template_path(e)).slides),1)

    def test_mirrors_keep_node_ids_endpoints_order_and_boundaries(self):
        for e in self.entries:
            if e['direction']!='lr' or e['kind']=='compare':continue
            source,plan=fixture([e]);before=plan.slides[0].contents.model_dump()
            mirror=plan.model_copy(deep=True);mirror.slides[0].layout_id=e['layout_id'][:-2]+'rl'
            self.assertEqual(self.codes(mirror.slides[0]),[])
            self.assertEqual(before,mirror.slides[0].contents.model_dump())
            deck=Presentation(self.deck(source,mirror))
            note=json.loads(deck.slides[0].notes_slide.notes_text_frame.text)
            self.assertEqual(note['relation_contents'],before)
            self.assertTrue(audit(self.base/'result.pptx',Theme())['ok'])
            if e['kind']=='pair_fan':
                mirrored=relation_registry()[mirror.slides[0].layout_id]
                for key,node in [('left_caption',0),('right_caption',2)]:
                    caption=mirrored['texts'][key]['bounds_emu'];box=mirrored['relation']['nodes'][node]['bounds_emu']
                    self.assertEqual(caption[0]<6000000,box[0]<6000000,'panel caption must mirror with its semantic nodes')

    def test_directional_routes_touch_only_their_declared_nodes(self):
        for e in self.entries:
            r=e['relation']
            for key,route in r.get('routes',{}).items():
                a,z=map(int,key.split('>'));x1,y1,x2,y2=route['points_emu']
                for i,x,y in ((a,x1,y1),(z,x2,y2)):
                    bx,by,w,h=r['nodes'][i]['bounds_emu']
                    self.assertTrue(bx-2<=x<=bx+w+2 and by-2<=y<=by+h+2,(e['layout_id'],key,i))
                    self.assertTrue(min(abs(x-bx),abs(x-bx-w),abs(y-by),abs(y-by-h))<=2)

    def test_rejects_unknown_duplicate_disconnected_and_unsupported_edges(self):
        source,plan=self.sample('rel_fan_4_lr');s=plan.slides[0]
        cases=[('RELATION_ENDPOINT',lambda n:setattr(n.edges[0],'target','missing')),
               ('RELATION_ID',lambda n:setattr(n.nodes[1],'id',n.nodes[0].id)),
               ('RELATION_DUPLICATE_EDGE',lambda n:n.edges.append(n.edges[0].model_copy(update={'id':'different'}))),
               ('RELATION_TOPOLOGY',lambda n:(setattr(n.edges[0],'source','actor_1'),setattr(n.edges[0],'target','actor_2'))),
               ('RELATION_DISCONNECTED',lambda n:setattr(n,'edges',n.edges[:2]))]
        for code,change in cases:
            trial=s.model_copy(deep=True);change(trial.contents.network);self.assertIn(code,self.codes(trial))

    def test_boundary_kinds_members_and_no_boundary_endpoints(self):
        source,plan=self.sample('rel_spoke_4_lr');s=plan.slides[0]
        trial=s.model_copy(deep=True);trial.contents.network.boundaries[0].members=['actor_1','actor_2']
        self.assertIn('RELATION_BOUNDARY',self.codes(trial))
        trial=s.model_copy(deep=True);trial.contents.network.edges[0].target='owner'
        self.assertIn('RELATION_ENDPOINT',self.codes(trial))
        trial=plan.model_copy(deep=True);trial.slides[0].contents.network.boundaries[0].kind='same_owner'
        self.assertEqual(self.codes(trial.slides[0]),[])
        p=self.deck(source,trial);self.assertTrue(audit(p,Theme())['ok'])
        self.assertIn('同管理：',next(s for s in Presentation(p).slides[0].shapes if s.name=='rel:boundary_label').text)

    def test_event_limits_self_processing_long_labels_and_empty_text(self):
        source,plan=self.sample('rel_sequence_4_8_lr');s=plan.slides[0]
        self.assertEqual(self.codes(s),[]);self.assertTrue(any(e.source==e.target for e in s.contents.network.edges))
        shorter=s.model_copy(update={'layout_id':'rel_sequence_4_4_lr'})
        self.assertIn('RELATION_EVENT_COUNT',self.codes(shorter))
        for value,code in [(' ', 'RELATION_EMPTY'),('長'*100,'RELATION_OVERFLOW')]:
            trial=s.model_copy(deep=True);trial.contents.network.edges[0].label.text=value;self.assertIn(code,self.codes(trial))
        data=plan.model_dump();data['slides'][0]['contents']['network']['nodes'].append(data['slides'][0]['contents']['network']['nodes'][0])
        self.assertIn('RELATION_NODE_COUNT',self.codes(Plan.model_validate(data).slides[0]))
        data['slides'][0]['contents']['network']['nodes'].append(data['slides'][0]['contents']['network']['nodes'][0])
        with self.assertRaises(ValidationError):Plan.model_validate(data)
        data=plan.model_dump();data['slides'][0]['contents']['network']['edges'][0]['x']=100
        with self.assertRaises(ValidationError):Plan.model_validate(data)

    def test_three_method_axes_equal_columns_and_outside_note(self):
        source,plan=self.sample('rel_compare_three');path=self.deck(source,plan)
        slide=Presentation(path).slides[0];table=next(s.table for s in slide.shapes if s.has_table)
        self.assertEqual([c.width for c in list(table.columns)[1:]],[table.columns[1].width]*3)
        c=plan.slides[0].contents.comparison
        for col,m in enumerate(c.methods,1):
            self.assertEqual(table.cell(0,col).text,m.label.text)
            for row,t in enumerate(m.values,1):self.assertEqual(table.cell(row,col).text,t.text)
        data=plan.model_dump();data['slides'][0]['contents']['comparison']['methods'].pop()
        with self.assertRaises(ValidationError):Plan.model_validate(data)

    def test_selection_preserves_graph_and_rejects_incompatible_topology(self):
        source,plan=self.sample('rel_sequence_4_8_lr');before=plan.model_dump()
        selected,report=select_variants(plan)
        self.assertEqual(before,plan.model_dump());self.assertEqual(selected.slides[0].contents,plan.slides[0].contents)
        self.assertTrue(report);self.assertTrue(all('_4_' not in x.split('sequence_4')[1] for x in report[0]['fitting']))

    def test_audit_detects_reversed_arrows_labels_containment_and_missing_edges(self):
        source,plan=self.sample('rel_spoke_4_lr');path=self.deck(source,plan)
        for tamper in ('direction','label','boundary','missing'):
            prs=Presentation(path);slide=prs.slides[0]
            if tamper=='direction':
                shape=next(s for s in slide.shapes if s.name.startswith('rel:edge_') and not s.name.startswith('rel:edge_label'))
                tail=shape._element.spPr.get_or_add_ln().find(qn('a:tailEnd'));tail.set('type','none')
            elif tamper=='label':next(s for s in slide.shapes if s.name=='rel:node_label_0').text='改変'
            elif tamper=='boundary':next(s for s in slide.shapes if s.name=='rel:boundary').width+=10000
            else:
                shape=next(s for s in slide.shapes if s.name.startswith('rel:edge_'));shape._element.getparent().remove(shape._element)
            out=self.base/'tamper.pptx';prs.save(out);self.assertFalse(audit(out,Theme())['ok'],tamper)

    def test_reader_hyperlinks_versions_notes_qa_and_safe_rejection(self):
        source,plan=self.sample('rel_pair_2_lr');path=self.deck(source,plan)
        slide=Presentation(path).slides[0];footer=next(s for s in slide.shapes if s.name=='meta:footer')
        self.assertIn('架空設計書',footer.text);self.assertIn('2026-10-07',footer.text);self.assertIn('v1',footer.text)
        self.assertEqual({r.hyperlink.address for p in footer.text_frame.paragraphs for r in p.runs if r.hyperlink.address},{'https://example.com/fictional-design'})
        self.assertNotIn('native editable','\n'.join(s.text for s in slide.shapes if s.has_text_frame))
        plan.display.mode='qa';qa=self.deck(source,plan,'qa.pptx');self.assertTrue(audit(qa,Theme())['ok'])
        self.assertEqual(next(s for s in Presentation(qa).slides[0].shapes if s.name=='meta:footer').text,'slide-layout-agent')
        plan.display.mode='reader';plan.display.citations[0].title='長'*100
        self.assertIn('DISPLAY_OVERFLOW',{i['code'] for i in validate(plan,source,self.base,Theme())['issues']})
        plan.display.citations[0].accessed_on='2026-02-31'
        self.assertIn('DISPLAY_DATE',{i['code'] for i in validate(plan,source,self.base,Theme())['issues']})
        data=plan.model_dump();data['display']['citations'][0]['url']='file:///private'
        with self.assertRaises(ValidationError):Plan.model_validate(data)

    def test_cli_discovers_relation_contract(self):
        result=subprocess.run([sys.executable,'-m','slide_agent','catalog','--collection','relations'],cwd=ROOT,capture_output=True,check=True)
        self.assertEqual(len(json.loads(result.stdout)),25)

    def test_sparse_pair_and_four_independent_actors(self):
        entries=[relation_registry()[key] for key in ('rel_pair_2_lr','rel_spoke_4_lr','rel_sequence_4_8_lr','rel_pair_fan_5_lr')]
        source,plan=fixture(entries,usage=True)
        self.assertEqual(len(plan.slides[0].contents.network.edges),1)
        self.assertEqual(plan.slides[2].contents.network.boundaries,[])
        self.assertEqual({(e.source,e.target) for e in plan.slides[1].contents.network.edges},{('actor_1','actor_0'),('actor_0','actor_2'),('actor_3','actor_0')})
        self.assertTrue(validate(plan,source,self.base,Theme())['ok'])
        result=audit(self.deck(source,plan),Theme());self.assertTrue(result['ok'],result['issues'])


if __name__=='__main__':unittest.main()
