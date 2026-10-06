import copy
import hashlib
import json
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

from PIL import Image
from pptx import Presentation
from pydantic import ValidationError

from scripts.make_demo import make_demo
from slide_agent.audit import audit
from slide_agent.cli import ingest, main, split_plan
from slide_agent.io import asset_path, fingerprint, load_plan, read_json
from slide_agent.models import Plan, Source, Theme
from slide_agent.render import picture, render
from slide_agent.validate import numbers, validate


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.source, self.plan = make_demo(self.root)
        self.theme = Theme()

    def tearDown(self):
        self.temp.cleanup()

    def report(self, plan=None, source=None, theme=None):
        return validate(plan or self.plan, source or self.source, self.root, theme or self.theme)

    def codes(self, **kwargs):
        return {i["code"] for i in self.report(**kwargs)["issues"]}

    def test_fictional_demo_all_eight_layouts(self):
        self.assertEqual(len({s.layout_id for s in self.plan.slides}), 8)
        self.assertEqual(self.report()["issues"], [])

    def test_codex_authored_plans_cover_both_input_modes(self):
        from slide_agent.io import load_source
        examples = Path(__file__).resolve().parents[1] / "examples"
        for stem in ("prose", "slides"):
            plan = load_plan(examples / f"{stem}.plan.json")
            source = load_source(examples / f"{stem}.source.json")
            report = validate(plan, source, examples, self.theme)
            self.assertEqual(report["issues"], [], stem)

    def test_ingest_prose_japanese_preserves_paragraphs_and_image(self):
        raw = "調査済みの文章です。\n\n参加者は12人。ただし暫定値です。"
        path = self.root / "input.txt"
        path.write_text(raw, encoding="utf-8")
        source = ingest(path, "prose", [f"photo={self.root / 'grid.png'}"], self.root / "source.json")
        self.assertEqual([s.text for s in source.segments], raw.split("\n\n"))
        self.assertEqual(source.images[0].path, "grid.png")

    def test_ingest_per_slide_preserves_groups(self):
        path = self.root / "slides.md"
        path.write_text("## 目的\n\n内容です。\n\n## 次の手順\n本文です。", encoding="utf-8")
        source = ingest(path, "slides", [], self.root / "source.json")
        self.assertEqual([s.group for s in source.segments], ["slide1", "slide1", "slide2", "slide2"])
        self.assertEqual([s.text for s in source.segments], ["目的", "内容です。", "次の手順", "本文です。"])

    def test_source_mutation_detected(self):
        source = self.source.model_copy(deep=True)
        source.segments[0].text += "改変"
        self.assertIn("SOURCE_CHANGED", self.codes(source=source))

    def test_unknown_layout_and_coordinates_rejected(self):
        data = self.plan.model_dump(mode="json")
        data["slides"][0]["layout_id"] = "run_python"
        with self.assertRaises(ValidationError):
            Plan.model_validate(data)
        data = self.plan.model_dump(mode="json")
        data["slides"][0]["x"] = 1
        with self.assertRaises(ValidationError):
            Plan.model_validate(data)

    def test_source_number_added_and_removed_are_errors(self):
        plan = self.plan.model_copy(deep=True)
        plan.slides[6].contents.series[0].values[0].value = 99
        self.assertIn("UNSUPPORTED_NUMBER", self.codes(plan=plan))
        self.assertIn("MISSING_NUMBER", self.codes(plan=plan))
        self.assertEqual(numbers("約１２人、1,200件、-3.5度、20%。"), {"12", "1.2E+3", "-3.5", "2E+1"})

    def test_qualifier_removal_triggers_partial_review(self):
        source = self.source.model_copy(deep=True)
        plan = self.plan.model_copy(deep=True)
        segment = source.segments[1]
        segment.text += "。ただし効果は未確認。"
        text = plan.slides[0].contents.items[0]
        text.refs[0].quote = segment.text
        plan.source_sha256 = fingerprint(source)
        self.assertIn("PARTIAL_SOURCE_REVIEW", self.codes(plan=plan, source=source))

    def test_paraphrase_warns_and_false_verbatim_rejected(self):
        plan = self.plan.model_copy(deep=True)
        plan.slides[0].title.text = "地域の情報共有"
        self.assertIn("NOT_VERBATIM", self.codes(plan=plan))
        plan.slides[0].title.mode = "paraphrase"
        self.assertIn("PARAPHRASE_REVIEW", self.codes(plan=plan))

    def test_missing_image_and_hash_change(self):
        source = self.source.model_copy(deep=True)
        source.images[0].path = "missing.png"
        plan = self.plan.model_copy(update={"source_sha256": fingerprint(source)})
        self.assertIn("IMAGE_INVALID", self.codes(plan=plan, source=source))
        (self.root / "grid.png").write_bytes(b"invalid")
        self.assertIn("IMAGE_INVALID", self.codes())

    def test_image_escape_rejected(self):
        asset = self.source.images[0].model_copy(update={"path": "../secret.png"})
        with self.assertRaises(ValueError):
            asset_path(asset, self.root)
        for path in ("https://example.com/x.png", "C:/private/x.png", "..\\x.png"):
            with self.assertRaises(ValueError):
                asset_path(asset.model_copy(update={"path": path}), self.root)

    def test_long_text_stops_without_shrink(self):
        plan = self.plan.model_copy(deep=True)
        plan.slides[1].contents.items[0].text = "長い本文です。" * 1000
        self.assertIn("TEXT_OVERFLOW", self.codes(plan=plan))
        with self.assertRaisesRegex(ValueError, "one item exceeds"):
            split_plan(plan, self.theme)

    def test_bullet_pagination_preserves_all_content_order_and_refs(self):
        plan = self.plan.model_copy(deep=True)
        original = plan.slides[1].contents.items * 3
        plan.slides[1].contents.items = original
        self.assertIn("ITEM_LIMIT", self.codes(plan=plan))
        result = split_plan(plan, self.theme)
        pages = [s for s in result.slides if s.origin == "section2"]
        self.assertGreater(len(pages), 1)
        self.assertEqual([t.model_dump() for s in pages for t in s.contents.items], [t.model_dump() for t in original])
        self.assertTrue(self.report(plan=result)["ok"])

    def test_table_row_overflow_and_split(self):
        plan = self.plan.model_copy(deep=True)
        original = plan.slides[5].contents.rows * 3
        plan.slides[5].contents.rows = original
        self.assertIn("ROW_LIMIT", self.codes(plan=plan))
        result = split_plan(plan, self.theme)
        pages = [s for s in result.slides if s.layout_id == "table"]
        self.assertEqual([row for s in pages for row in s.contents.rows], original)
        self.assertTrue(self.report(plan=result)["ok"])

    def test_table_dimensions_and_chart_values_checked(self):
        data = self.plan.model_dump(mode="json")
        data["slides"][5]["contents"]["rows"][0].pop()
        with self.assertRaises(ValidationError):
            Plan.model_validate(data)
        data = self.plan.model_dump(mode="json")
        data["slides"][6]["contents"]["series"][0]["values"].pop()
        with self.assertRaises(ValidationError):
            Plan.model_validate(data)

    def test_omission_is_explicit_and_cannot_conflict(self):
        data = self.plan.model_dump(mode="json")
        data["omissions"] = [{"source_id": self.source.segments[0].id, "reason": "テスト"}]
        self.assertIn("OMISSION_CONFLICT", self.codes(plan=Plan.model_validate(data)))
        plan = self.plan.model_copy(deep=True)
        plan.slides.pop(0)
        self.assertIn("UNCOVERED_SOURCE", self.codes(plan=plan))

    def test_minimum_sizes_and_lead_policy(self):
        with self.assertRaises(ValidationError):
            Theme(body_pt=12)
        plan = self.plan.model_copy(deep=True)
        plan.slides[0].lead = plan.slides[0].contents.items[0]
        self.assertIn("LEAD_DISABLED", self.codes(plan=plan))
        self.assertNotIn("LEAD_DISABLED", self.codes(plan=plan, theme=Theme(lead_mode=True)))

    def test_image_fit_and_crop_preserve_aspect(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        fit = picture(slide, self.root / "grid.png", 1, 1, 4, 4, "fit")
        self.assertAlmostEqual(fit.width / fit.height, 1.5, places=5)
        crop = picture(slide, self.root / "grid.png", 6, 1, 4, 4, "crop")
        self.assertAlmostEqual(crop.width / crop.height, 1, places=5)
        self.assertAlmostEqual((1 - crop.crop_left - crop.crop_right) * 1.5, 1, places=4)
        self.assertEqual(crop.crop_top, 0)

    def test_native_ooxml_and_roundtrip_editability(self):
        out = self.root / "test.pptx"
        render(self.plan, self.source, self.root, self.theme, out)
        result = audit(out, self.theme)
        self.assertTrue(result["ok"], result)
        self.assertEqual(result["counts"]["tables"], 1)
        self.assertEqual(result["counts"]["charts"], 1)
        self.assertEqual(result["counts"]["embedded_workbooks"], 1)
        prs = Presentation(out)
        title = next(s for s in prs.slides[0].shapes if s.name == "title")
        title.text = "編集後の見出し"
        table = next(s.table for s in prs.slides[5].shapes if s.has_table)
        table.cell(1, 1).text = "編集後の表"
        from pptx.chart.data import CategoryChartData
        chart = next(s.chart for s in prs.slides[6].shapes if s.has_chart)
        data = CategoryChartData()
        data.categories = ["編集前", "編集後"]
        data.add_series("件数", [1, 2])
        chart.replace_data(data)
        edited = self.root / "edited.pptx"
        prs.save(edited)
        reopened = Presentation(edited)
        self.assertIn("編集後の見出し", [s.text for s in reopened.slides[0].shapes if s.has_text_frame])
        self.assertEqual(next(s.table for s in reopened.slides[5].shapes if s.has_table).cell(1, 1).text, "編集後の表")
        self.assertEqual(list(next(s.chart for s in reopened.slides[6].shapes if s.has_chart).series[0].values), [1, 2])
        with ZipFile(out) as z:
            self.assertIn(b"typeface=\"Meiryo\"", z.read("ppt/slides/slide1.xml"))
            self.assertIn(b"source_sha256", z.read("ppt/notesSlides/notesSlide1.xml"))

    def test_deterministic_slide_xml(self):
        a, b = self.root / "a.pptx", self.root / "b.pptx"
        render(self.plan, self.source, self.root, self.theme, a)
        render(self.plan, self.source, self.root, self.theme, b)
        with ZipFile(a) as za, ZipFile(b) as zb:
            names = [n for n in za.namelist() if n.startswith(("ppt/slides/", "ppt/charts/")) and n.endswith(".xml")]
            self.assertTrue(all(za.read(n) == zb.read(n) for n in names))

    def test_duplicate_json_key_rejected(self):
        path = self.root / "duplicate.json"
        path.write_text('{"version":"1","version":"2"}', encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "duplicate JSON key"):
            read_json(path)

    def test_cli_end_to_end_and_no_overwrite(self):
        out = self.root / "cli.pptx"
        args = ["render", str(self.root / "demo.plan.json"), "--source", str(self.root / "demo.source.json"), "--out", str(out)]
        self.assertEqual(main(args), 0)
        self.assertTrue(out.exists())
        self.assertTrue(out.with_suffix(".validation.json").exists())
        self.assertEqual(main(args), 2)


if __name__ == "__main__":
    unittest.main()
